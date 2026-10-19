#!/usr/bin/env python3
"""Copia los resultados de plantilla_cribado.xlsx a manuscrito/datos_busqueda.tex.

Lee las hojas «Busqueda», «PRISMA» y «Kappa» de la plantilla (los valores que
Excel calculó la última vez que se guardó el archivo) y reemplaza los \\pend{...}
correspondientes en datos_busqueda.tex. Lo que aún esté vacío en la plantilla
se deja como \\pend{...} y se lista al final.

Uso (desde la carpeta RSL):
    python suplementos/actualizar_datos_busqueda.py

Importante: abrir la plantilla en Excel y GUARDAR antes de ejecutar el script,
para que las fórmulas tengan su valor calculado.
"""
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "suplementos" / "plantilla_cribado.xlsx"
DATOS = RAIZ / "manuscrito" / "datos_busqueda.tex"

MESES_EN = ["January", "February", "March", "April", "May", "June", "July",
            "August", "September", "October", "November", "December"]


def fecha(v):
    """Devuelve (dd/mm/aaaa, 'd Month aaaa') a partir de una fecha o texto."""
    if isinstance(v, (datetime, date)):
        d = v
    else:
        d = datetime.strptime(str(v).strip(), "%d/%m/%Y")
    return d.strftime("%d/%m/%Y"), f"{d.day} {MESES_EN[d.month - 1]} {d.year}"


def numero(v):
    if isinstance(v, float) and not v.is_integer():
        return f"{v:.2f}".replace(".", ",")  # decimal con coma, como el resto del texto
    return str(int(v))


def leer_valores():
    wb = load_workbook(PLANTILLA, data_only=True)
    val = {}
    # Búsqueda: fechas y resultados por base
    busq = wb["Busqueda"]
    claves = {"Scopus": "Scopus", "Web of Science": "WoS", "IEEE Xplore": "IEEE",
              "ScienceDirect": "ScienceDirect", "ACM Digital Library": "ACM"}
    fechas = []
    for fila in busq.iter_rows(min_row=2, max_row=6, values_only=True):
        base, f, _, n = fila[:4]
        k = claves.get(base)
        if not k:
            continue
        if n not in (None, ""):
            val[f"N{k}"] = numero(n)
        if f not in (None, ""):
            es, _ = fecha(f)
            val[f"F{k}"] = es
            fechas.append(f if isinstance(f, (datetime, date)) else datetime.strptime(str(f).strip(), "%d/%m/%Y"))
    if fechas:
        ultima = max(fechas)
        val["FechaBusqueda"], val["FechaBusquedaEN"] = fecha(ultima)
    # PRISMA: columna B = valor, columna C = macro
    for fila in wb["PRISMA"].iter_rows(min_row=2, values_only=True):
        if len(fila) < 3 or not fila[2] or not str(fila[2]).startswith("\\"):
            continue
        if fila[1] in (None, "") or isinstance(fila[1], str):
            continue
        val[str(fila[2]).lstrip("\\")] = numero(fila[1])
    # Kappa: buscar las filas cuyo rótulo menciona la macro
    for fila in wb["Kappa"].iter_rows(values_only=True):
        if fila[0] and "KAPPA" in str(fila[0]) and fila[1] not in (None, ""):
            macro = str(fila[0]).split("\\")[-1].strip()
            val[macro] = numero(float(fila[1]))
    return val


def main():
    val = leer_valores()
    tex = DATOS.read_text(encoding="utf-8")
    cambiados = []
    for macro, v in val.items():
        patron = re.compile(r"(\\newcommand\{\\" + macro + r"\}\{)\\pend\{[^}]*\}(\})")
        tex, n = patron.subn(lambda m: m.group(1) + v + m.group(2), tex)
        if n:
            cambiados.append(f"{macro} = {v}")
    DATOS.write_text(tex, encoding="utf-8")
    print("Actualizados:" if cambiados else "No se actualizó nada (¿guardaste la plantilla en Excel?).")
    for c in cambiados:
        print("  " + c)
    faltan = re.findall(r"\\newcommand\{\\(\w+)\}(?:\[\d\])?\{\\pend", tex)
    if faltan:
        print("Siguen pendientes en datos_busqueda.tex:", ", ".join(faltan))


if __name__ == "__main__":
    main()
