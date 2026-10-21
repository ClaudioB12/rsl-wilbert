#!/usr/bin/env python3
"""PASO 12. Pasa a TA.xlsx lo que decidiste en la lista de revisión (LISTA_REVISION_TA.xlsx).

En la lista llenas, para cada artículo que revisas:
  MI DECISIÓN   «De acuerdo» (acepta la propuesta), Incluir, Duda o Excluir
  MI CRITERIO   obligatorio si excluyes (CE-1, CE-2, CE-3, CE-4, CE-5, CI-2, CI-3)
  MI COMENTARIO tu motivo en una frase (si lo dejas vacío se usa el motivo de la IA)
Las filas sin MI DECISIÓN no se tocan.

Uso (desde la carpeta RSL):
    python cribado/scripts/12_aplicar_lista_revision.py              # revisa y aplica
    python cribado/scripts/12_aplicar_lista_revision.py --probar     # solo muestra lo que haría
    python cribado/scripts/12_aplicar_lista_revision.py --archivo cribado/revision/OTRA_LISTA.xlsx
Acepta las columnas «MI COMENTARIO» o «MI COMENTARIO / MOTIVO». En título y resumen CE-6 no
se usa (es de texto completo); la falta de acceso abierto es CI-2.
Luego: python cribado/scripts/09_generar_script_parsifal.py y pega el script en Parsif.al.
El comentario queda sin «IA (verificar):», porque ya lo verificaste tú.
"""
import argparse
import re
import sys
from collections import Counter

from openpyxl import load_workbook

from comun import REVISION, asegurar_utf8_consola
from hojas import COL_CRITERIO, COL_DECISION

LISTA = REVISION / "LISTA_REVISION_TA.xlsx"
CRITERIOS = {"CE-1", "CE-2", "CE-3", "CE-4", "CE-5", "CI-2", "CI-3"}
RE_PREFIJO = re.compile(r"^\s*(\[[^\]]*\]\s*)?(IA \(verificar\):\s*)?")


def texto(v):
    return "" if v is None else str(v).strip()


def leer_lista(ruta, propuesta):
    """propuesta: {id: (decisión, criterio, motivo)} tomada de TA.xlsx (columnas «… IA»)."""
    wb = load_workbook(ruta, data_only=True)
    filas, errores = {}, []
    for ws in wb.worksheets:
        cab = [texto(c.value) for c in ws[1]]
        if "MI DECISIÓN" not in cab or "ID" not in cab:
            continue
        col_com = next((n for n in ("MI COMENTARIO", "MI COMENTARIO / MOTIVO") if n in cab), None)
        for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            g = lambda k: texto(fila[cab.index(k)]) if k in cab and cab.index(k) < len(fila) else ""
            i, mi = g("ID"), g("MI DECISIÓN")
            if not i or not mi:
                continue
            if i not in propuesta:
                errores.append(f"{ws.title} fila {n}: el ID «{i}» no está en TA.xlsx")
                continue
            p_dec, p_crit, p_mot = propuesta[i]
            if mi.lower().startswith("de acuerdo"):
                dec, crit = p_dec, p_crit
            elif mi in ("Incluir", "Duda", "Excluir"):
                dec, crit = mi, g("MI CRITERIO").upper()
            else:
                errores.append(f"{ws.title} fila {n} ({i}): MI DECISIÓN «{mi}» no válida")
                continue
            if dec == "Excluir":
                if crit == "CE-6":
                    errores.append(f"{ws.title} fila {n} ({i}): CE-6 solo se usa a texto completo; si no es de acceso "
                                   "abierto el criterio es CI-2")
                    continue
                if crit not in CRITERIOS:
                    errores.append(f"{ws.title} fila {n} ({i}): Excluir necesita MI CRITERIO válido (tiene «{crit}»)")
                    continue
            else:
                crit = ""
            comentario = (g(col_com) if col_com else "") or RE_PREFIJO.sub("", p_mot)
            filas[i] = {"decision": dec, "criterio": crit, "comentario": RE_PREFIJO.sub("", comentario),
                        "cambio": (dec, crit) != (p_dec, p_crit if p_dec == "Excluir" else ""), "hoja": ws.title}
    return filas, errores


def main():
    asegurar_utf8_consola()
    ap = argparse.ArgumentParser()
    ap.add_argument("--probar", action="store_true", help="solo mostrar, sin cambiar TA.xlsx")
    ap.add_argument("--archivo", default=str(LISTA), help="lista de revisión llenada (por defecto LISTA_REVISION_TA.xlsx)")
    a = ap.parse_args()
    from pathlib import Path
    if not Path(a.archivo).exists():
        sys.exit(f"No encuentro {a.archivo}.")
    from hojas import leer_hoja
    propuesta = {i: (v.get("Propuesta IA", ""), v.get("Criterio IA", ""), v.get("Comentario IA", ""))
                 for i, v in leer_hoja(REVISION / "TA.xlsx").items()}
    filas, errores = leer_lista(a.archivo, propuesta)
    if errores:
        print(f"Hay {len(errores)} filas para corregir en la lista (no se aplicó nada):")
        for e in errores[:30]:
            print("   " + e)
        sys.exit(1)
    if not filas:
        sys.exit("No hay ninguna fila con MI DECISIÓN. Llena la lista y vuelve a ejecutar.")

    por_hoja = Counter(f["hoja"] for f in filas.values())
    cambios = [i for i, f in filas.items() if f["cambio"]]
    print("Filas revisadas por hoja: " + ", ".join(f"{h}: {n}" for h, n in sorted(por_hoja.items())))
    print(f"Decisiones que cambias respecto de la propuesta de la IA: {len(cambios)}")
    print("Resultado: " + ", ".join(f"{k} {n}" for k, n in sorted(Counter(
        f["decision"] + (" " + f["criterio"] if f["criterio"] else "") for f in filas.values()).items())))
    for i in cambios[:40]:
        f = filas[i]
        print(f"   {i}: → {f['decision']} {f['criterio']}".rstrip())
    if a.probar:
        print("(Modo prueba: no se cambió TA.xlsx.)")
        return

    ruta = REVISION / "TA.xlsx"
    wb = load_workbook(ruta)
    ws = wb["Cribado"]
    cab = [c.value for c in ws[1]]
    cd, cc, cm = cab.index(COL_DECISION) + 1, cab.index(COL_CRITERIO) + 1, cab.index("Comentario") + 1
    aplicadas = 0
    for fila in range(2, ws.max_row + 1):
        f = filas.get(str(ws.cell(row=fila, column=1).value))
        if not f:
            continue
        ws.cell(row=fila, column=cd).value = f["decision"]
        ws.cell(row=fila, column=cc).value = f["criterio"] or None
        ws.cell(row=fila, column=cm).value = f["comentario"]
        aplicadas += 1
    wb.save(ruta)
    print(f"TA.xlsx actualizado: {aplicadas} filas.")
    print("Ahora: python cribado\\scripts\\09_generar_script_parsifal.py  y pega cargar_en_parsifal.js en Parsif.al.")


if __name__ == "__main__":
    main()
