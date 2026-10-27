#!/usr/bin/env python3
"""PASO 15. Pasa al cribado la lectura a texto completo (REVISION_TEXTO_COMPLETO.xlsx).

Dos pasos:
  A) python cribado/scripts/15_aplicar_texto_completo.py --normalizar
     Limpia el Excel: lleva cada celda de las listas (tipo de dato, entorno, etiquetas, paradigma,
     familia, híbrido, categoría, partición, línea base no DL) a una de las opciones de la lista
     (el texto original queda como comentario de la celda), pone el código del criterio de exclusión,
     marca «No recuperado» las filas sin texto completo y aplica las correcciones de la revisión
     (AJUSTES, en amarillo). Las filas que hay que rehacer quedan en rojo (PENDIENTES). Todo se
     resume en la hoja «4_Observaciones». Guarda una copia del original en *_antes_de_normalizar.xlsx.
  B) python cribado/scripts/15_aplicar_texto_completo.py
     Valida el Excel y, si no falta nada, escribe:
       - cribado/salida/decision_TC.json        (lo usa 06_resultados_prisma.py)
       - cribado/salida/15_extraccion_TC.json   (extracción y calidad de los incluidos)
       - cribado/salida/parsifal/cargar_TC_en_parsifal.js  (decisiones a texto completo en Parsif.al)
     Con --probar solo muestra el resumen y los errores.
"""
import argparse
import importlib.util
import json
import re
import shutil
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

from comun import REVISION, SALIDA, asegurar_utf8_consola, cargar_json, guardar_json, leer_csv
from hojas import AJUSTE, AMARILLO, AZUL, BLANCO_NEGRITA, ROJO

SCRIPTS = Path(__file__).resolve().parent
LIBRO = REVISION / "REVISION_TEXTO_COMPLETO.xlsx"
ORDEN_TC = ["CE-1", "CE-5", "CI-3", "CI-2", "CE-2", "CE-3", "CE-4", "CI-6", "CE-6"]
DECISIONES = ["Incluir", "Excluir", "No recuperado"]


def _modulo(nombre):
    spec = importlib.util.spec_from_file_location(nombre.replace(".py", ""), SCRIPTS / nombre)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


PLANTILLA = _modulo("14_plantilla_texto_completo.py")
OPCIONES = {c[0]: c[2] for c in PLANTILLA.H2 if c[2]}
CAMPOS_H2 = [c[0] for c in PLANTILLA.H2]
ITEMS_QA = [q for q, _ in PLANTILLA.QA]

# Correcciones de la revisión del autor con el asistente: (ID, hoja, columna) → (valor, motivo).
AJUSTES = {
    ("R0244", 1, "DECISIÓN TC"): ("Incluir", "«INCLUIDO» no es una opción de la lista."),
    ("R0190", 1, "CRITERIO"): ("CE-3", "Es tráfico vehicular (no CE-2): solo pronostica el flujo con MAE, no detecta anomalías."),
    ("R0249", 1, "CRITERIO"): ("CI-6", "Sí trata la detección de anomalías, pero solo de forma exploratoria y sin métricas de detección: es CI-6, no CE-3."),
    ("R0528", 1, "CRITERIO"): ("CI-6", "Detecta accidentes con DL, pero no reporta métricas de detección: es CI-6, no CE-3."),
    ("R1166", 1, "CRITERIO"): ("CI-6", "Es un sistema de detección de incidentes sin métricas de detección: es CI-6, no CE-3."),
    ("R0614", 1, "CRITERIO"): ("CE-3", "Predice la gravedad de accidentes ya registrados; CE-3 va antes que CE-4 en el orden fijo."),
    ("R0604", 1, "DECISIÓN TC"): ("Excluir", "Monitoreo del conductor dentro de la cabina (fatiga, bostezos, teléfono): CE-2 lo excluye expresamente."),
    ("R0604", 1, "CRITERIO"): ("CE-2", "Monitoreo del conductor dentro de la cabina."),
    ("R0929", 1, "DECISIÓN TC"): ("Excluir", "La CNN solo reconoce si el conductor sostiene el teléfono (imagen de cabina): CE-2 lo excluye expresamente."),
    ("R0929", 1, "CRITERIO"): ("CE-2", "Monitoreo del conductor dentro de la cabina."),
    ("R0389", 1, "CRITERIO"): ("CI-6", "Trata anomalías de tráfico pero no reporta métricas de detección: es CI-6 (igual que R0249, R0528 y R1166)."),
    ("R0403", 3, "QA5 Código y datos"): (1, "Código con DOI (ETS) y datasets públicos (DAD, CCD, A3D): ambos públicos."),
    ("R0032", 2, "Tipo de dato"): ("Multimodal", "Acelerómetro, giroscopio y GPS del teléfono (UAH-DriveSet)."),
    ("R0403", 2, "Partición"): ("Prueba independiente", "Particiones oficiales de DAD y CCD."),
    ("R0192", 1, "DECISIÓN TC"): ("Excluir", "Releído el PDF: la detección solo se valida de forma visual (face validation); "
                                              "la única métrica cuantitativa (RMSE, Tabla 1) es de imputación."),
    ("R0192", 1, "CRITERIO"): ("CI-6", "Sin evaluación cuantitativa de la detección (igual que R0249 y R0389)."),
    ("R0192", 1, "EVIDENCIA (cita y página)"): (
        '"we employ a face validation approach. In this approach, we visually examine the pace time series of various zones" '
        '(p. 10, sección IV-E); "Table 1 presents the mean and standard deviation of RMSE across 10 trials of inference-phase '
        'imputation" (p. 9, sección IV-C).', "Cita del PDF."),
    ("R0192", 1, "Comentario"): (
        "[Excluir CI-6] SI-DMSeq (encoder-decoder LSTM con atención, Manhattan, huracán Sandy y tormenta de 2010) detecta "
        "disrupciones con una carta de control sobre el error de reconstrucción, pero no las evalúa contra etiquetas: no hay "
        "precision, recall, F1, AUC ni tasa de detección; la única métrica numérica (RMSE) es de imputación.", "Motivo."),
    ("R0167", 2, "Paradigma"): ("No supervisado", "Releído el PDF: MGTGCN se entrena para predecir y reconstruir trayectorias "
                                                "(ADE/FDE, Tabla 1, p. 12), sin etiquetas de anomalía; las conductas anómalas se "
                                                "añaden después para evaluar: «adding a set of disruptive driving behaviors» (p. 12)."),
    ("R0029", 1, "DECISIÓN TC"): ("No recuperado", "El autor verificó que no hay texto completo en acceso abierto "
                                                 "(la copia de SSRN no está disponible)."),
    ("R0029", 1, "CRITERIO"): (None, "No recuperado no lleva criterio."),
    ("R0029", 1, "Comentario"): ("[No recuperado] Sin texto completo de acceso abierto.", "Coherente con las demás no recuperadas."),
    # Auditoría de coherencia de los incluidos (mismos criterios que en decisiones previas del cribado).
    ("R1014", 1, "DECISIÓN TC"): ("Excluir", "DeepWalk/Node2vec no es Deep Learning; la anomalía la deciden SVDD y reglas."),
    ("R1014", 1, "CRITERIO"): ("CE-4", "Sin componente profunda entrenada que decida la anomalía."),
    ("R1014", 1, "EVIDENCIA (cita y página)"): (
        "DeepWalk genera caminatas aleatorias y aprende embeddings con Skip-Gram; SVDD determina las regiones anómalas "
        "(sección 3, pp. 3-5); umbral de coincidencia de trayectorias anómalas 0,463 (sección 4.2, p. 6).", "Hoja 2 del autor."),
    ("R1014", 1, "Comentario"): ("[Excluir CE-4] El detector es DeepWalk/Node2vec (embedding superficial, no DL) más SVDD, "
                                 "Apriori y umbrales.", "Motivo."),
    ("R0666", 1, "DECISIÓN TC"): ("Excluir", "YOLOv5s/DeepSORT solo siguen vehículos; la anomalía se decide con umbrales "
                                             "manuales (igual que R0037, R0235 y R1213, excluidos por CE-4)."),
    ("R0666", 1, "CRITERIO"): ("CE-4", "La decisión de anomalía no la toma una red profunda."),
    ("R0666", 1, "EVIDENCIA (cita y página)"): (
        "Reglas explícitas: sentido contrario si Num/f ≥ 0,2; velocidad anómala si la media de 10 fotogramas supera n o baja "
        "de m; estacionamiento si la posición no cambia en cinco fotogramas (sección 3.3.5, pp. 10-11; Tabla 1, p. 8).",
        "Hoja 2 del autor."),
    ("R0666", 1, "Comentario"): ("[Excluir CE-4] El DL detecta, sigue y reidentifica vehículos; los comportamientos "
                                 "peligrosos se deciden con umbrales manuales.", "Motivo."),
    ("R0032", 1, "DECISIÓN TC"): ("Excluir", "Perfilado del conductor por viaje (igual que R0321, R0660 y R1046, CE-3)."),
    ("R0032", 1, "CRITERIO"): ("CE-3", "Clasifica estilos/perfiles de conducción, no anomalías del tráfico."),
    ("R0032", 1, "EVIDENCIA (cita y página)"): (
        "UAH-DriveSet, seis conductores, tres comportamientos (normal, somnoliento y agresivo); clasificación supervisada "
        "en una de las tres categorías (p. 3, secciones 3.1-3.2).", "Hoja 1 y 2 del autor."),
    ("R0032", 1, "Comentario"): ("[Excluir CE-3] Compara ML, DL y un ensamble PSO para perfilar al conductor "
                                 "(normal/somnoliento/agresivo); no detecta eventos anómalos del tráfico.", "Motivo."),
    ("R0072", 1, "DECISIÓN TC"): ("Excluir", "Mismo modelo (MHGNN-CPLE), mismos autores y mismos datos que R0102; "
                                             "se conserva R0102 (Scientific Reports, versión más completa)."),
    ("R0072", 1, "CRITERIO"): ("CE-1", "Versión menos completa del mismo estudio."),
    ("R0072", 1, "EVIDENCIA (cita y página)"): (
        "Propone MHGNN-CPLE evaluado en METR-LA y PEMS-BAY con generalización METR-LA→PEMS-BAY (pp. 21-23, Tablas 5, 7 y 8); "
        "R0102 (Duan, Sun y Liu, 2025) presenta el mismo marco con las mismas bases.", "Hojas 1 y 2 del autor."),
    ("R0072", 1, "Comentario"): ("[Excluir CE-1] Publicación paralela de R0102 (Duan et al., 2025); se incluye R0102 para no "
                                 "contar dos veces el mismo modelo.", "Motivo."),
    ("R0185", 1, "DECISIÓN TC"): ("Excluir", "Clasifica niveles de riesgo de interacción, no eventos anómalos "
                                             "(igual que R0635, excluido por CE-3)."),
    ("R0185", 1, "CRITERIO"): ("CE-3", "Estimación de riesgo, sin detección de anomalías."),
    ("R0185", 1, "EVIDENCIA (cita y página)"): (
        "Clasificación del riesgo de seguimiento y cambio de carril en niveles bajo, moderado y alto: Accuracy 83 % y 76 % "
        "(p. 6, Tablas 3 y 4).", "Hoja 1 del autor."),
    ("R0185", 1, "Comentario"): ("[Excluir CE-3] CNN-LSTM-Attention estima el nivel de riesgo de las interacciones en túneles; "
                                 "no detecta ni anticipa eventos anómalos.", "Motivo."),
    ("R0190", 1, "Comentario"): ("[Excluir CE-3] Pronostica el flujo con Deep LSTM y Mixture Deep LSTM (MAE) en horas punta y "
                                 "tras accidentes; los accidentes son una condición de entrada, no eventos que se detecten.", "Coherente con CE-3."),
    ("R0249", 1, "Comentario"): ("[Excluir CI-6] VAE para imputación, reducción de dimensión y pronóstico; la detección de anomalías "
                                 "es exploratoria (espacio latente) y no se evalúa con etiquetas ni métricas de detección.", "Coherente con CI-6."),
    ("R0528", 1, "Comentario"): ("[Excluir CI-6] YOLO-R, DeepSORT y GAN para detectar accidentes en video, sin tabla de resultados "
                                 "de detección (precision, recall, F1 o AUC) frente a accidentes etiquetados.", "Coherente con CI-6."),
    ("R1166", 1, "Comentario"): ("[Excluir CI-6] Plataforma de videovigilancia con detección y seguimiento de vehículos; muestra "
                                 "ejemplos de incidentes, pero no reporta métricas de detección.", "Coherente con CI-6."),
    ("R0614", 1, "Comentario"): ("[Excluir CE-3] Clasifica la gravedad de accidentes ya registrados (Random Forest, XGBoost, "
                                 "LightGBM); no detecta incidentes y tampoco usa DL.", "Coherente con CE-3."),
    ("R0604", 1, "Comentario"): ("[Excluir CE-2] YOLOv8n reconoce fatiga, bostezos, uso del teléfono y tabaquismo del conductor de camión en imágenes de cabina: es monitoreo del conductor.", "Comentario conciso para el anexo."),
    ("R0929", 1, "Comentario"): ("[Excluir CE-2] La CNN clasifica si el conductor sostiene el teléfono en imágenes de cabina; el módulo de accidentes depende del acelerómetro y no se evalúa con DL.", "Comentario conciso para el anexo."),
    ("R0389", 1, "Comentario"): ("[Excluir CI-6] Propone un enfoque de ciencia de datos con DL para anomalías de velocidad, pero no reporta métricas de detección (precision, recall, F1 o AUC) ni etiquetas de referencia.", "Comentario conciso para el anexo."),
    ("R0585", 1, "Comentario"): ("[Excluir CE-5] El texto completo disponible es una tesis de maestría; no se encontró una versión revisada por pares de acceso abierto.", "Comentario conciso para el anexo."),
    ("R0618", 1, "Comentario"): ("[Excluir CE-5] Preprint de arXiv (2505.02613) sin versión revisada por pares verificable.", "Comentario conciso para el anexo."),
    ("R0622", 1, "Comentario"): ("[Excluir CE-5] Preprint de arXiv (2303.05000) sin versión revisada por pares verificable.", "Comentario conciso para el anexo."),
    ("R0635", 1, "Comentario"): ("[Excluir CE-3] GPT-2-LoRA, GAT y LSTM predicen el riesgo de colisión en eventos peligrosos preseleccionados (TTC < 3 s) de highD; no detecta anomalías sobre tráfico no preseleccionado.", "Comentario conciso para el anexo."),
    ("R0646", 1, "Comentario"): ("[Excluir CE-3] LSTM difusa que pronostica variables de conducción y un nivel de riesgo calculado con reglas difusas; solo reporta MAE, MSE y RMSE.", "Comentario conciso para el anexo."),
    ("R0986", 1, "Comentario"): ("[Excluir CE-3] LSTM con IoT para pronosticar el flujo de autopista; menciona eventos anormales, pero no los etiqueta ni evalúa su detección.", "Comentario conciso para el anexo."),
    ("R0024", 3, "QA8 Limitaciones"): (0, "La hoja 2 dice «No reportadas»: la cita es trabajo futuro, no una limitación."),
    ("R0024", 3, "QA8 evidencia"): ("No reportado (solo trabajo futuro, p. 10)", "Coherente con la hoja 2."),
    ("R0121", 2, "Entorno"): ("Urbano", "Dos cruces peatonales de calle (35 mph) en Florida."),
    ("R0052", 2, "Entorno"): ("Ambos", "Intersecciones y rampas/puentes de carretera."),
    ("R0119", 2, "Entorno"): ("Ambos", "Intersecciones y autopistas."),
    ("R0501", 2, "Entorno"): ("Ambos", "Vías urbanas e interurbanas."),
    ("R0175", 2, "Entorno"): ("No reportado", "No especifica el tipo de vía."),
    ("R0598", 2, "Entorno"): ("No reportado", "Pista de pruebas; no es urbano ni autopista."),
    ("R1000", 2, "Entorno"): ("No reportado", "No indica qué escenas de NGSIM usa."),
    ("R1010", 2, "Entorno"): ("No reportado", "Evaluado con autos de juguete en escenario controlado."),
    ("R1138", 2, "Entorno"): ("No reportado", "No desagrega urbano/autopista."),
    ("R1185", 2, "Entorno"): ("No reportado", "No clasifica las imágenes por tipo de vía."),
    ("R0097", 2, "Entorno"): ("No reportado", "Ubicación y tipo de vía no reportados."),
    ("R1100", 2, "Entorno"): ("No reportado", "DoTA/CCD sin clasificación por tipo de vía."),
    ("R1148", 2, "Categoría de anomalía"): ("Evento/otro", "Objetos caídos sobre la calzada."),
    ("R1014", 2, "Familia DL principal"): ("Otra", "DeepWalk/Node2vec no es una GNN."),
    ("R1148", 2, "Partición"): ("No reportada", "Solo dice que entrena y valida 100 épocas."),
    ("R0119", 2, "Partición"): ("Aleatoria", "Partición estratificada sin separación temporal."),
    ("R0185", 2, "Partición"): ("No reportada", "No describe la partición."),
    ("R0167", 2, "Partición"): ("No reportada", "«No se detallan particiones»."),
    ("R0501", 2, "Partición"): ("Prueba independiente", "Conjunto externo ECVD solo para prueba."),
    ("R0080", 2, "Partición"): ("Prueba independiente", "Incidentes distintos en entrenamiento (23) y prueba (22)."),
    ("R0379", 2, "Partición"): ("Prueba independiente", "Evalúa en una grabación highD no vista."),
    ("R1092", 2, "Partición"): ("Aleatoria", "División 80/20 de videos sin otro criterio."),
    ("R0943", 2, "Partición"): ("No reportada", "Sin procedimiento de selección ni conjunto de prueba."),
    ("R1013", 2, "Partición"): ("Prueba independiente", "Conjunto de prueba oficial de TADS."),
    ("R1082", 2, "Partición"): ("Prueba independiente", "Conjunto de prueba de MAAD (con ajuste de hiperparámetros en parte de él)."),
    ("R0658", 2, "Partición"): ("Prueba independiente", "Particiones oficiales de los benchmarks."),
    ("R0664", 2, "Partición"): ("Prueba independiente", "Videos de prueba separados."),
    ("R0666", 2, "Partición"): ("No reportada", "No detalla la partición."),
    ("R0596", 2, "Partición"): ("No reportada", "Partición insuficientemente reportada."),
    ("R0601", 2, "Partición"): ("Prueba independiente", "Conjunto de prueba de AI City Challenge."),
    ("R0537", 2, "Partición"): ("No reportada", "Sin proporciones ni protocolo de separación."),
    ("R1214", 2, "Partición"): ("No reportada", "No informa proporciones ni criterio de división."),
    ("R1202", 2, "Partición"): ("Prueba independiente", "Separación por sesiones de conducción."),
}
# Filas que hay que volver a hacer con su propio PDF (no se pueden corregir sin leerlo).
PENDIENTES = {
}
# Avisos que no cambian nada, pero que conviene mirar.
AVISOS = {
    "R0986": "El texto analizado corresponde a Systems and Soft Computing (DOI 10.1016/j.sasc.2025.200254), pero el registro tiene "
             "DOI 10.1002/spy2.70136: confirma que el PDF es el del registro.",
    "R0025": "Partición «No reportada» en la hoja 2 pero QA2 = 1 en la hoja 3: revisa uno de los dos.",
    "R0051": "Partición aleatoria (411/35/148) pero QA2 = 1: según la escala sería 0.5.",
}


def sin_tildes(t):
    t = unicodedata.normalize("NFKD", str(t or ""))
    return "".join(c for c in t if not unicodedata.combining(c)).lower()


REGLAS = {
    "Tipo de dato": [("Multimodal", r"multimodal"), ("Video", r"video|imagen|cctv|camara|fotograma|dashcam|dron|uav|aere"),
                     ("Trayectorias GPS", r"trayector|gps|telemetr"),
                     ("Sensores/detectores", r"sensor|detector|radar|espira|bucle|loop|series? tempor"),
                     ("Movilidad agregada", r"movilidad|agregad|amap|origen-destino|taxi")],
    "Entorno": [("Ambos", r"ambos|urbano.{0,40}(autopista|autovia|freeway|highway)|(autopista|autovia).{0,40}urban"),
                ("Autopista", r"autopista|autovia|freeway|highway|expressway|interestatal|interstate|tunel|carretera"),
                ("Urbano", r"urban|ciudad|intersecci|calle"), ("No reportado", r"no (se )?(report|especific|indic)")],
    "Origen de etiquetas": [("Inyectadas/sintéticas", r"inyect|sintetic|simulad|carla|sumo"), ("Pseudo-etiquetas", r"pseudo|reglas"),
                            ("Anotación manual", r"anotac|manual|anotad|expert"),
                            ("Reales/registros", r"reales|registro|oficial|dataset|conjunto|existentes|preexistentes"),
                            ("Sin etiquetas", r"sin etiquetas|no existen etiquetas|no requiere etiquetas"),
                            ("No reportado", r"no report|no detall|no especific")],
    "Paradigma": [("Semi-supervisado", r"semi|debilmente"), ("Auto-supervisado", r"auto-?supervis|self-?supervis|contrastiv"),
                  ("No supervisado", r"no supervis|unsupervis"), ("Supervisado", r"supervis")],
    "Familia DL principal": [("Autoencoder/VAE", r"autoencoder|\bvae\b|encoder-decoder|codificador"),
                             ("Generativo (GAN/difusión)", r"\bgan\b|generativ|difusi"),
                             ("Grafos (GNN)", r"gnn|grafo|graph|gcn|\bgat\b"),
                             ("Atención/Transformer", r"transformer|atencion|attention|\bvit\b|swin|clip"),
                             ("Recurrente (RNN/LSTM/GRU)", r"recurrent|lstm|gru|\brnn\b"),
                             ("Convolucional (CNN)", r"cnn|convolu|yolo|resnet|efficientnet|mobilenet"),
                             ("Otra", r"dnn|red neuronal|mlp|reinforcement|refuerzo|elm|otra")],
    "Categoría de anomalía": [("Accidentes/incidentes", r"accident|incident|colisi|choque|crash"),
                              ("Fallo de sensor/datos", r"fallo de sensor|sensor fail|datos faltantes|lecturas erron"),
                              ("Comportamiento de conducción", r"conducci|driving|maniobra|conductor"),
                              ("Trayectoria anómala", r"trayector"), ("Congestión atípica", r"congesti|flujo")],
    "Partición": [("No reportada", r"^\s*no (report|especific)"), ("Temporal", r"^temporal|separacion temporal|cronol|dias completos|periodo"),
                  ("Por ubicación", r"ubicaci|por conductor|por vehiculo"),
                  ("Prueba independiente", r"independ|estandar|oficial|externo|por incidentes|conjuntos separados"),
                  ("Aleatoria", r"aleator|estratific|validacion cruzada|cross-?valid|random|\d+\s*[:/%]|\d+ entrenamiento"),
                  ("No reportada", r"no (se )?(report|especific|document|detall)")],
}


def _primero(t, reglas):
    mejor = None
    for op, rx in reglas:
        m = re.search(rx, t)
        if m and (mejor is None or m.start() < mejor[0]):
            mejor = (m.start(), op)
    return mejor[1] if mejor else None


def normalizar(campo, valor):
    """Devuelve la opción de la lista que corresponde al texto (o None si no se reconoce)."""
    crudo = str(valor or "").strip()
    if crudo in OPCIONES[campo]:
        return crudo
    t = sin_tildes(crudo)
    if not t:
        return None
    if campo in ("¿Híbrido?", "¿Línea base no DL?"):
        return "Sí" if re.match(r"\s*(si|yes)\b", t) else "No" if re.match(r"\s*no\b", t) else None
    op = _primero(re.split(r"[;:.(—]", t)[0], REGLAS[campo]) or _primero(t, REGLAS[campo])
    return op or ("Evento/otro" if campo == "Categoría de anomalía" else None)


def codigo_criterio(texto):
    """Primer criterio del orden fijo que aparece al inicio del texto («CE-4: …», «CE-3 (incumple CI-6)»)."""
    codigos = re.findall(r"C[IE]-\d", str(texto or "").upper())
    return codigos[0] if codigos else ""


def numero(v):
    try:
        x = float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None
    return x if x in (0, 0.5, 1) else None


def hojas_por_id(wb):
    out = {}
    for n, nombre in ((1, "1_TextoCompleto"), (2, "2_Extraccion"), (3, "3_Calidad")):
        ws = wb[nombre]
        cab = [c.value for c in ws[1]]
        out[n] = (ws, cab, {str(ws.cell(row=f, column=1).value): f for f in range(3, ws.max_row + 1)
                            if ws.cell(row=f, column=1).value})
    return out


def comentar(celda, texto):
    if texto and str(texto).strip():
        celda.comment = Comment(f"Texto original:\n{str(texto)[:1500]}", "Normalización")


def normalizar_libro(ruta):
    copia = ruta.with_name(ruta.stem + "_antes_de_normalizar.xlsx")
    if not copia.exists():
        shutil.copy(ruta, copia)
    wb = load_workbook(ruta)
    H = hojas_por_id(wb)
    obs = []
    ws1, cab1, f1 = H[1]
    cd, cc, cm = cab1.index("DECISIÓN TC") + 1, cab1.index("CRITERIO") + 1, cab1.index("Comentario") + 1
    for i, f in f1.items():
        d = ws1.cell(row=f, column=cd)
        crit = ws1.cell(row=f, column=cc)
        txt = str(d.value or "").strip()
        if not txt and not ws1.cell(row=f, column=cm).value:
            d.value = "No recuperado"
            ws1.cell(row=f, column=cm).value = "[No recuperado] Sin texto completo de acceso abierto."
            obs.append((i, "1", "No recuperado", "Sin PDF: cuenta en PRISMA como «informe no recuperado», no como excluido."))
        elif txt.upper().startswith("INCL"):
            d.value = "Incluir"
        elif txt.upper().startswith("EXCL"):
            d.value = "Excluir"
        if d.value == "Incluir" and crit.value:
            comentar(crit, crit.value)
            crit.value = None
        elif d.value == "Excluir" and crit.value and crit.value not in ORDEN_TC:
            comentar(crit, crit.value)
            crit.value = codigo_criterio(crit.value)
    ws1.data_validations.dataValidation = [v for v in ws1.data_validations.dataValidation
                                           if str(v.sqref).split(":")[0][0] != "E"]
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONES) + '"', allow_blank=True)
    ws1.add_data_validation(dv)
    dv.add(f"E3:E{ws1.max_row}")

    ws2, cab2, f2 = H[2]
    for i, f in f2.items():
        if ws1.cell(row=f1[i], column=cd).value != "Incluir":
            continue
        for campo in OPCIONES:
            c = ws2.cell(row=f, column=cab2.index(campo) + 1)
            if c.value in OPCIONES[campo]:
                continue
            nuevo = normalizar(campo, c.value)
            comentar(c, c.value)
            c.value = nuevo
    ws3, cab3, f3 = H[3]
    for i, f in f3.items():
        for q in ITEMS_QA:
            c = ws3.cell(row=f, column=cab3.index(q) + 1)
            if numero(c.value) is not None:
                c.value = numero(c.value)

    for (i, hoja, col), (valor, motivo) in AJUSTES.items():
        ws, cab, filas = H[hoja]
        c = ws.cell(row=filas[i], column=cab.index(col) + 1)
        if c.value != valor:
            if c.value not in (None, ""):
                c.comment = Comment(f"Antes: {str(c.value)[:500]}\nMotivo: {motivo}", "Revisión")
            c.value = valor
        c.fill = AMARILLO
        obs.append((i, str(hoja), f"Corregido «{col}» → {valor}", motivo))
    for i, (hs, motivo) in PENDIENTES.items():
        for h in hs.split(","):
            ws, _, filas = H[int(h)]
            for c in ws[filas[i]]:
                c.fill = ROJO
        obs.append((i, hs, "REHACER", motivo))
    for i, motivo in AVISOS.items():
        obs.append((i, "-", "Revisar", motivo))

    if "4_Observaciones" in wb.sheetnames:
        del wb["4_Observaciones"]
    ws = wb.create_sheet("4_Observaciones", 1)
    ws.append(["ID", "Hoja", "Qué", "Motivo"])
    for c, w in zip(ws[1], (8, 7, 40, 110)):
        c.fill, c.font = AZUL, BLANCO_NEGRITA
        ws.column_dimensions[c.column_letter].width = w
    orden = {"REHACER": 0, "Revisar": 1}
    for o in sorted(obs, key=lambda o: (orden.get(o[2], 2), o[0])):
        ws.append(list(o))
        for c in ws[ws.max_row]:
            c.alignment = AJUSTE
            if o[2] == "REHACER":
                c.fill = ROJO
    ws.freeze_panes = "A2"
    wb.save(ruta)
    print(f"Normalizado {ruta.name}: {len(AJUSTES)} correcciones (amarillo), {len(PENDIENTES)} filas a rehacer (rojo). "
          f"Ver la hoja «4_Observaciones». Copia del original: {copia.name}")


def leer(ruta):
    wb = load_workbook(ruta, data_only=True)
    H = hojas_por_id(wb)
    reg = {r["id"]: r for r in leer_csv(SALIDA / "02_registros_unicos.csv")}
    ta = cargar_json(SALIDA / "decision_TA.json") or {}
    pasan = {i for i, x in ta.items() if x["decision"] == "Incluir"}
    errores, estudios = [], {}
    fila = lambda h, i: dict(zip(H[h][1], [c.value for c in H[h][0][H[h][2][i]]])) if i in H[h][2] else {}
    for i in sorted(pasan):
        a, b, c = fila(1, i), fila(2, i), fila(3, i)
        if not a:
            errores.append(f"{i}: no está en la hoja 1")
            continue
        d = a.get("DECISIÓN TC")
        e = {"id": i, "titulo": reg[i]["titulo"], "anio": reg[i]["anio"], "doi": reg[i]["doi"], "decision": d,
             "criterio": a.get("CRITERIO") or "", "evidencia": str(a.get("EVIDENCIA (cita y página)") or "").strip(),
             "nota": str(a.get("Comentario") or "").strip()}
        if i in PENDIENTES and d == "Incluir":
            errores.append(f"{i}: pendiente de rehacer (hojas {PENDIENTES[i][0]}): {PENDIENTES[i][1]}")
        if d not in DECISIONES:
            errores.append(f"{i}: DECISIÓN TC «{d}» no válida")
        elif d == "Excluir":
            if e["criterio"] not in ORDEN_TC:
                errores.append(f"{i}: Excluir necesita un CRITERIO válido (tiene «{e['criterio']}»)")
            if not e["evidencia"]:
                errores.append(f"{i}: Excluir necesita EVIDENCIA con cita y página")
        elif d == "Incluir":
            e["extraccion"] = {k: (b.get(k) if not isinstance(b.get(k), str) else b.get(k).strip()) for k in CAMPOS_H2}
            for k in CAMPOS_H2:
                v = e["extraccion"][k]
                if v in (None, "") or str(v).upper() == "EXCLUIDO":
                    errores.append(f"{i}: hoja 2, «{k}» vacío")
                elif k in OPCIONES and v not in OPCIONES[k]:
                    errores.append(f"{i}: hoja 2, «{k}» = «{str(v)[:40]}» no es una opción de la lista")
            e["calidad"] = {}
            for q in ITEMS_QA:
                p = numero(c.get(q))
                if p is None:
                    errores.append(f"{i}: hoja 3, {q} = «{c.get(q)}» (debe ser 1, 0.5 o 0)")
                e["calidad"][q[:3]] = {"p": p, "ev": str(c.get(f"{q[:3]} evidencia") or "").strip()}
            e["calidad_total"] = sum(x["p"] or 0 for x in e["calidad"].values())
        estudios[i] = e
    return estudios, errores


def script_parsifal(estudios):
    gen = _modulo("09_generar_script_parsifal.py")
    datos = {}
    for i, e in estudios.items():
        if e["decision"] == "Incluir":
            datos[i] = ["A", "", f"[Incluir TC] {e['nota']}"[:4000], gen.norm(e["titulo"])]
        elif e["decision"] == "Excluir":
            datos[i] = ["R", e["criterio"], f"[Excluir TC {e['criterio']}] {e['nota']} Evidencia: {e['evidencia']}"[:4000],
                        gen.norm(e["titulo"])]
        else:
            datos[i] = ["R", "", "[No recuperado] Sin texto completo de acceso abierto.", gen.norm(e["titulo"])]
    salida = SALIDA / "parsifal" / "cargar_TC_en_parsifal.js"
    salida.write_text(gen.PLANTILLA.replace("__DATOS__", json.dumps(datos, ensure_ascii=False))
                      .replace("09_generar_script_parsifal.py", "15_aplicar_texto_completo.py"), encoding="utf-8")
    return salida


def main():
    asegurar_utf8_consola()
    ap = argparse.ArgumentParser()
    ap.add_argument("--archivo", default=str(LIBRO))
    ap.add_argument("--normalizar", action="store_true", help="limpiar el Excel (paso A)")
    ap.add_argument("--probar", action="store_true", help="solo validar y mostrar el resumen")
    a = ap.parse_args()
    ruta = Path(a.archivo)
    if not ruta.exists():
        sys.exit(f"No encuentro {ruta}.")
    if a.normalizar:
        normalizar_libro(ruta)
        return
    estudios, errores = leer(ruta)
    c = Counter(e["decision"] for e in estudios.values())
    crit = Counter(e["criterio"] for e in estudios.values() if e["decision"] == "Excluir")
    print(f"Estudios a texto completo: {len(estudios)} → " + ", ".join(f"{k} {n}" for k, n in sorted(c.items(), key=str)))
    print("Exclusiones por criterio: " + ", ".join(f"{k} {n}" for k, n in sorted(crit.items())))
    if errores:
        print(f"\nFaltan {len(errores)} cosas por corregir en {ruta.name} (no se escribió nada):")
        for e in errores[:40]:
            print("   " + e)
        sys.exit(1)
    if a.probar:
        print("(Modo prueba: no se escribió nada.)")
        return
    guardar_json(SALIDA / "decision_TC.json", {i: {**{k: e[k] for k in ("decision", "criterio", "evidencia", "nota")},
                                                    "decision_original": e["decision"]}  # para 04b (re-cribado)
                                                for i, e in estudios.items()})
    guardar_json(SALIDA / "15_extraccion_TC.json", {i: e for i, e in estudios.items() if e["decision"] == "Incluir"})
    js = script_parsifal(estudios)
    print(f"Escrito decision_TC.json, 15_extraccion_TC.json y {js.name}.")
    print("Ahora: python cribado/scripts/06_resultados_prisma.py   y pega cargar_TC_en_parsifal.js en Parsif.al.")


if __name__ == "__main__":
    main()
