#!/usr/bin/env python3
"""PASO 7. Evaluación de calidad (QA1-QA8) por dos revisores independientes.

Uso (desde la carpeta RSL):
  A) python cribado/scripts/07_calidad.py --preparar
       Crea cribado/revision/calidad_R1.xlsx y calidad_R2.xlsx (vacías, con las definiciones).
  B) python cribado/scripts/07_calidad.py --consolidar
       Calcula el kappa (3 categorías: 1 / 0.5 / 0) y crea salida/consolidado_calidad.xlsx,
       con el consenso prellenado donde coinciden.
  C) python cribado/scripts/07_calidad.py --aplicar
       Lleva el consenso a suplementos/extraccion_individual/*.json. Luego ejecuten
       python suplementos/generar_conteos.py para regenerar las tablas.
"""
import argparse
import json
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from comun import REVISION, SALIDA, SUPLEMENTOS, asegurar_utf8_consola, cargar_json, guardar_json, interpretar_kappa, kappa_cohen
from hojas import AJUSTE, AMARILLO, AZUL, BLANCO_NEGRITA, ROJO, _lista

ITEMS = {
    "QA1": ("Procedencia de etiquetas", "1: registros reales o anotación verificada; 0.5: mixtas, automáticas o verificación parcial; 0: solo inyectadas o no reportado"),
    "QA2": ("Partición sin fuga", "1: partición temporal o prueba independiente descrita; 0.5: partición aleatoria o descripción incompleta; 0: no reportada"),
    "QA3": ("Líneas base comparables", "1: dos o más líneas base con el mismo protocolo, incluida una del estado del arte o no DL; 0.5: una sola o solo ablación; 0: ninguna"),
    "QA4": ("Variabilidad", "1: varias ejecuciones con dispersión, intervalos o pruebas de significancia; 0.5: parcial; 0: una sola ejecución"),
    "QA5": ("Código y datos", "1: ambos públicos; 0.5: uno de los dos; 0: ninguno"),
    "QA6": ("Más de un escenario", "1: dos o más datasets, ciudades o vías independientes; 0.5: varios subconjuntos de un dataset; 0: un único escenario"),
    "QA7": ("Protocolo explícito", "1: reporta unidad de evaluación y criterio de umbral o decisión; 0.5: solo uno; 0: ninguno"),
    "QA8": ("Limitaciones", "1: sección o párrafo explícito; 0.5: mención breve; 0: ninguna"),
}
VALORES = ["1", "0.5", "0"]


def estudios():
    return json.loads((SUPLEMENTOS / "matriz_extraccion.json").read_text(encoding="utf-8"))


def preparar(forzar):
    for rev in ("R1", "R2"):
        ruta = REVISION / f"calidad_{rev}.xlsx"
        if ruta.exists() and not forzar:
            sys.exit(f"Ya existe {ruta} (usa --forzar si estás seguro).")
        wb = Workbook()
        ws = wb.active
        ws.title = "Calidad"
        cab = ["ID", "Estudio"] + [f"{q} {ITEMS[q][0]}" for q in ITEMS] + ["Justificación breve de los puntajes 0 o 0.5"]
        ws.append(cab)
        for i, c in enumerate(ws[1], 1):
            c.fill, c.font, c.alignment = AZUL, BLANCO_NEGRITA, AJUSTE
            ws.column_dimensions[c.column_letter].width = 7 if i == 1 else 40 if i in (2, len(cab)) else 13
        ws.row_dimensions[1].height = 45
        for e in estudios():
            ws.append([e["id"], f"{e['primer_autor']} et al. ({e['anio']}). {e['titulo']}"] + [None] * 9)
            for c in ws[ws.max_row]:
                c.alignment = AJUSTE
        _lista(ws, f"C2:J{ws.max_row}", VALORES)
        ws.freeze_panes = "C2"
        d = wb.create_sheet("Definiciones")
        d.column_dimensions["A"].width = 8
        d.column_dimensions["B"].width = 28
        d.column_dimensions["C"].width = 110
        d.append(["Ítem", "Criterio", "Puntaje"])
        for q, (n, txt) in ITEMS.items():
            d.append([q, n, txt])
        d.append([])
        d.append(["", "Revisor", rev])
        d.append(["", "Regla", "Puntuar leyendo el PDF completo, sin mirar la hoja del otro revisor ni la matriz de extracción."])
        for fila in d.iter_rows():
            for c in fila:
                c.alignment = AJUSTE
        wb.save(ruta)
        print(f"Creado {ruta}")


def leer(rev):
    ws = load_workbook(REVISION / f"calidad_{rev}.xlsx", data_only=True)["Calidad"]
    out = {}
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if fila[0]:
            out[fila[0]] = {q: ("" if fila[2 + j] is None else str(fila[2 + j]).replace(",", ".").rstrip("0").rstrip(".") or "0")
                            for j, q in enumerate(ITEMS)}
            out[fila[0]]["just"] = fila[10] or ""
    return out


def consolidar(forzar):
    destino = SALIDA / "consolidado_calidad.xlsx"
    if destino.exists() and not forzar:
        sys.exit(f"Ya existe {destino} (usa --forzar si estás seguro).")
    r1, r2 = leer("R1"), leer("R2")
    pares = [(r1[i][q], r2[i][q]) for i in r1 for q in ITEMS if i in r2]
    k, n, po = kappa_cohen(pares)
    wb = Workbook()
    ws = wb.active
    ws.title = "Consenso"
    cab = ["ID"] + [f"{q} R1" for q in ITEMS] + [f"{q} R2" for q in ITEMS] + [f"{q} CONSENSO" for q in ITEMS] + ["Nota"]
    ws.append(cab)
    for c in ws[1]:
        c.fill, c.font, c.alignment = AZUL, BLANCO_NEGRITA, AJUSTE
    for i in r1:
        a, b = r1[i], r2.get(i, {})
        cons = [a[q] if a[q] == b.get(q) and a[q] != "" else None for q in ITEMS]
        ws.append([i] + [a[q] for q in ITEMS] + [b.get(q, "") for q in ITEMS] + cons + [None])
        for j, q in enumerate(ITEMS):
            if cons[j] is None:
                ws.cell(row=ws.max_row, column=18 + j).fill = ROJO
    _lista(ws, f"R2:Y{ws.max_row}", VALORES)
    ks = wb.create_sheet("Kappa", 0)
    for a_, b_ in [("Pares ítem×estudio", n), ("Acuerdo observado", f"{po:.1%}" if po is not None else "-"),
                   ("Kappa de Cohen", f"{k:.2f}" if k is not None else "-"), ("Interpretación", interpretar_kappa(k))]:
        ks.append([a_, b_])
        ks.cell(row=ks.max_row, column=1).font = Font(bold=True)
    ks.column_dimensions["A"].width = 28
    wb.save(destino)
    kap = cargar_json(SALIDA / "kappa.json", {})
    kap["calidad"] = {"kappa": f"{k:.2f}" if k is not None else None, "n": n, "acuerdo": f"{po:.1%}" if po is not None else None}
    guardar_json(SALIDA / "kappa.json", kap)
    print(f"Kappa calidad: {k:.2f} ({interpretar_kappa(k)}), n = {n}. Completen las celdas rojas en {destino} y luego --aplicar.")


def aplicar():
    ws = load_workbook(SALIDA / "consolidado_calidad.xlsx", data_only=True)["Consenso"]
    cur = json.loads((SUPLEMENTOS / "curacion.json").read_text(encoding="utf-8"))
    clave = {c["id"]: c["clave_extraccion"] for c in cur["incluidos"]}
    faltan, cambios = [], 0
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila[0]:
            continue
        cons = fila[17:25]
        if any(v in (None, "") for v in cons):
            faltan.append(fila[0])
            continue
        ruta = SUPLEMENTOS / "extraccion_individual" / f"{clave[fila[0]]}.json"
        e = json.loads(ruta.read_text(encoding="utf-8"))
        for q, v in zip(ITEMS, cons):
            nuevo = float(str(v).replace(",", "."))
            if e["calidad"][q]["p"] != nuevo:
                e["calidad"][q]["ev"] = (e["calidad"][q]["ev"] + f" [Consenso de revisores: {nuevo}; antes {e['calidad'][q]['p']}]")
                e["calidad"][q]["p"] = nuevo
                cambios += 1
        ruta.write_text(json.dumps(e, ensure_ascii=False, indent=1), encoding="utf-8")
    if faltan:
        sys.exit(f"Faltan puntajes de consenso en: {', '.join(faltan)}")
    print(f"Consenso aplicado ({cambios} puntajes cambiaron). Ahora ejecuten: python suplementos/generar_conteos.py")


def main():
    asegurar_utf8_consola()
    ap = argparse.ArgumentParser()
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--preparar", action="store_true")
    g.add_argument("--consolidar", action="store_true")
    g.add_argument("--aplicar", action="store_true")
    ap.add_argument("--forzar", action="store_true")
    a = ap.parse_args()
    preparar(a.forzar) if a.preparar else consolidar(a.forzar) if a.consolidar else aplicar()


if __name__ == "__main__":
    main()
