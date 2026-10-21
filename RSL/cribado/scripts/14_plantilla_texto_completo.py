#!/usr/bin/env python3
"""PASO 14. Crea la plantilla en Excel para la lectura a texto completo de los estudios que pasaron
el título y resumen: decisión (hoja 1), extracción de datos (hoja 2) y calidad QA1-QA8 (hoja 3).

Uso (desde la carpeta RSL):   python cribado/scripts/14_plantilla_texto_completo.py [--forzar]
Salida: cribado/revision/REVISION_TEXTO_COMPLETO.xlsx

Las tres hojas tienen las mismas filas (una por estudio, con el ID R0001…). Las hojas 2 y 3 solo se
llenan para los estudios que quedan en «Incluir» en la hoja 1.
"""
import argparse
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation

from comun import REVISION, SALIDA, asegurar_utf8_consola, cargar_json, leer_csv
from hojas import AJUSTE, AMARILLO, AZUL, BLANCO_NEGRITA, VERDE_OSC

FIJAS = [("ID", 8), ("Título", 45), ("Año", 6), ("PDF", 16)]

# (encabezado, ancho, opciones de la lista desplegable o None, ayuda)
H1 = [
    ("DECISIÓN TC", 12, ["Incluir", "Excluir"], "Incluir o Excluir tras leer el texto completo"),
    ("CRITERIO", 10, ["CE-2", "CE-3", "CE-4", "CE-5", "CI-6", "CE-6", "CE-1", "CI-2", "CI-3"],
     "Solo si excluyes: el primer criterio que no cumple (orden: CE-5, CE-2, CE-3, CE-4, CI-6, CE-6)"),
    ("EVIDENCIA (cita y página)", 50, None, "Frase copiada del PDF y su página, p. ej. «…we forecast traffic flow…» (p. 4)"),
    ("Comentario", 35, None, "Tu motivo en una frase"),
]
H2 = [
    ("Tipo de dato", 16, ["Sensores/detectores", "Trayectorias GPS", "Movilidad agregada", "Video", "Multimodal"], "RQ1/RQ4"),
    ("Entorno", 12, ["Urbano", "Autopista", "Ambos", "No reportado"], "RQ4"),
    ("Datasets (nombre, país)", 30, None, "p. ej. PeMS (EE. UU.); UCF-Crime"),
    ("Origen de etiquetas", 16, ["Reales/registros", "Anotación manual", "Inyectadas/sintéticas", "Pseudo-etiquetas", "Sin etiquetas", "No reportado"], "RQ2"),
    ("Tasa de anomalías", 14, None, "Proporción de anomalías en la evaluación, o «No reportada»"),
    ("Paradigma", 14, ["Supervisado", "Semi-supervisado", "No supervisado", "Auto-supervisado"], "RQ1"),
    ("Familia DL principal", 16, ["Recurrente (RNN/LSTM/GRU)", "Convolucional (CNN)", "Grafos (GNN)", "Atención/Transformer", "Autoencoder/VAE", "Generativo (GAN/difusión)", "Otra"], "RQ1"),
    ("Arquitectura (nombre)", 25, None, "p. ej. ST-GCN + LSTM-AE"),
    ("¿Híbrido?", 9, ["Sí", "No"], "Combina dos o más familias"),
    ("Categoría de anomalía", 18, ["Accidentes/incidentes", "Congestión atípica", "Trayectoria anómala", "Comportamiento de conducción", "Fallo de sensor/datos", "Evento/otro"], "RQ3"),
    ("Métricas de detección", 25, None, "p. ej. F1, AUC-ROC, DR, FAR"),
    ("Resultado principal", 25, None, "p. ej. F1 = 0,94 en PeMS"),
    ("Unidad de evaluación", 16, None, "Punto, ventana, evento, trayectoria, fotograma, video…"),
    ("Criterio de umbral", 20, None, "Cómo decide anomalía (umbral fijo, percentil, clasificador…)"),
    ("Partición", 16, ["Temporal", "Aleatoria", "Por ubicación", "Prueba independiente", "No reportada"], "QA2"),
    ("Líneas base (cuántas y cuáles)", 30, None, "RQ6"),
    ("¿Línea base no DL?", 10, ["Sí", "No"], "RQ6"),
    ("Ventaja vs. mejor línea base", 18, None, "p. ej. +3,2 puntos de F1"),
    ("Limitaciones que reportan", 30, None, "RQ5"),
    ("Código público (URL)", 20, None, "Enlace o «No»"),
]
QA = [
    ("QA1 Etiquetas", "1: reales o anotación verificada; 0.5: mixtas o automáticas; 0: inyectadas o no reportado"),
    ("QA2 Partición", "1: temporal o prueba independiente; 0.5: aleatoria o incompleta; 0: no reportada"),
    ("QA3 Líneas base", "1: dos o más con el mismo protocolo, incluida una del estado del arte o no DL; 0.5: una o solo ablación; 0: ninguna"),
    ("QA4 Variabilidad", "1: varias ejecuciones con dispersión o significancia; 0.5: parcial; 0: una ejecución"),
    ("QA5 Código y datos", "1: ambos públicos; 0.5: uno; 0: ninguno"),
    ("QA6 Escenarios", "1: dos o más datasets/ciudades/vías; 0.5: subconjuntos de uno; 0: uno solo"),
    ("QA7 Protocolo", "1: unidad de evaluación y umbral; 0.5: solo uno; 0: ninguno"),
    ("QA8 Limitaciones", "1: sección o párrafo; 0.5: mención breve; 0: ninguna"),
]
H3 = [x for q, ayuda in QA for x in ((q, 9, ["1", "0.5", "0"], ayuda), (f"{q[:3]} evidencia", 28, None, "Frase y página que justifica el puntaje"))]


def hoja(wb, titulo, columnas, filas, total=False):
    ws = wb.create_sheet(titulo)
    cab = FIJAS + [(c[0], c[1]) for c in columnas] + ([("TOTAL QA (0-8)", 10)] if total else [])
    ws.append([c[0] for c in cab])
    ws.append([""] * len(FIJAS) + [c[3] for c in columnas] + ([""] if total else []))
    for i, (c, (_, ancho)) in enumerate(zip(ws[1], cab)):
        c.fill = AZUL if i < len(FIJAS) else VERDE_OSC
        c.font, c.alignment = BLANCO_NEGRITA, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = ancho
    for c in ws[2]:
        c.alignment, c.fill = AJUSTE, AMARILLO
    ws.row_dimensions[2].height = 60
    for f in filas:
        ws.append([f["id"], f["titulo"], f["anio"], f"cribado\\pdf\\{f['id']}.pdf"])
    n = ws.max_row
    for j, col in enumerate(columnas, len(FIJAS) + 1):
        letra = ws.cell(row=1, column=j).column_letter
        if col[2]:
            dv = DataValidation(type="list", formula1='"' + ",".join(col[2]) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letra}3:{letra}{n}")
    if total:
        letras = [ws.cell(row=1, column=j).column_letter for j in range(len(FIJAS) + 1, len(FIJAS) + len(columnas) + 1, 2)]
        for r in range(3, n + 1):
            celdas = ",".join(f"{L}{r}" for L in letras)
            ws.cell(row=r, column=len(cab)).value = f'=IF(COUNT({celdas})=0,"",SUM({celdas}))'
    for fila in ws.iter_rows(min_row=3, max_row=n):
        for c in fila:
            c.alignment = AJUSTE
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(cab)).column_letter}{n}"
    return ws


def main():
    asegurar_utf8_consola()
    ap = argparse.ArgumentParser()
    ap.add_argument("--forzar", action="store_true")
    a = ap.parse_args()
    ruta = REVISION / "REVISION_TEXTO_COMPLETO.xlsx"
    if ruta.exists() and not a.forzar:
        sys.exit(f"Ya existe {ruta}. No se sobrescribe (usa --forzar si estás seguro).")
    decision = cargar_json(SALIDA / "decision_TA.json")
    filas = [r for r in leer_csv(SALIDA / "02_registros_unicos.csv") if decision.get(r["id"], {}).get("decision") == "Incluir"]
    wb = Workbook()
    ins = wb.active
    ins.title = "0_Instrucciones"
    for linea in [
        "LECTURA A TEXTO COMPLETO — un estudio por fila, el mismo ID en las tres hojas.",
        "",
        "Hoja 1_TextoCompleto: lee el PDF y decide Incluir o Excluir. Si excluyes: CRITERIO + EVIDENCIA con cita y página.",
        "   Orden de los criterios: CE-5 (survey, preprint, resumen extendido) → CE-2 (no es tráfico vehicular) →",
        "   CE-3 (solo pronóstico, no detecta anomalías) → CE-4 (el DL no es el método principal) →",
        "   CI-6 (no evalúa cuantitativamente la DETECCIÓN; solo otra tarea) → CE-6 (no describe modelo, datos o resultados).",
        "Hoja 2_Extraccion: solo para los que quedan en Incluir. Usa las listas desplegables cuando las haya.",
        "Hoja 3_Calidad: solo para los incluidos. Cada ítem 1 / 0.5 / 0 y la frase con la página que lo justifica.",
        "   La calidad NO excluye estudios: solo se reporta.",
        "La fila amarilla (fila 2) explica qué poner en cada columna.",
        "Cuando termines, guarda el archivo en cribado\\revision\\ y súbelo con git (o pásalo por el chat).",
    ]:
        ins.append([linea])
    ins.column_dimensions["A"].width = 130
    hoja(wb, "1_TextoCompleto", H1, filas)
    hoja(wb, "2_Extraccion", H2, filas)
    hoja(wb, "3_Calidad", H3, filas, total=True)
    wb.save(ruta)
    print(f"Creado {ruta}: {len(filas)} estudios en 3 hojas (texto completo, extracción, calidad).")


if __name__ == "__main__":
    main()
