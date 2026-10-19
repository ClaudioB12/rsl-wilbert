#!/usr/bin/env python3
"""PASO 5b. El asistente lee los PDF y propone la decisión a texto completo, con citas y páginas.

Requisitos: pip install pypdf ; los PDF guardados como cribado/pdf/<ID>.pdf
Uso (desde la carpeta RSL):   python cribado/scripts/05b_leer_pdfs.py

Para cada artículo de TC.xlsx:
  - Extrae el texto página por página.
  - Busca evidencia de: tipo de documento (survey/preprint), Deep Learning, anomalía/detección,
    tráfico vial y MÉTRICAS DE DETECCIÓN (precision, recall, F1, AUC, tasa de detección, falsas
    alarmas...), y copia frases cortas con su página.
  - Escribe en TC.xlsx las columnas «SUGERENCIA», «Motivo» y «Evidencia sugerida: ...».
    NO toca tus columnas de decisión. Se puede ejecutar varias veces (por ejemplo, cuando
    descargues más PDF).
  - Marca «Texto completo = Sí» si encontró el PDF y la celda estaba vacía.
Salida adicional: cribado/salida/05b_evidencias.csv
"""
import logging
import re
import sys
import unicodedata
import warnings

from openpyxl import load_workbook

from comun import (CRIBADO, RE_ANOMALIA, RE_DL, RE_NO_VIAL, RE_VIAL, REVISION, SALIDA, asegurar_utf8_consola,
                   escribir_csv)
from hojas import COL_TEXTO

RE_METRICA = re.compile(r"\b(precision|recall|f1[- ]?score|f[- ]?measure|f1\b|auc|auroc|aupr|roc|pr[- ]auc|average precision|"
                        r"accuracy|detection rate|false alarm|false positive rate|\bfpr\b|\bfar\b|\btpr\b|"
                        r"mean time to detect|mttd|true positive|confusion matrix)", re.I)
# Métrica seguida de un valor numérico (p. ej., «F1 = 0.94», «AUC of 0.912», «accuracy of 97.01%»).
RE_METRICA_NUM = re.compile(r"\b(precision|recall|f1[- ]?score|f[- ]?measure|f1|auc(?:-roc)?|auroc|aupr|pr[- ]auc|"
                            r"accuracy|detection rate|false alarm rate|false positive rate|fpr|far|tpr|mttd)\b"
                            r"[^.\d]{0,30}?(\d{1,3}(?:[.,]\d+)?\s?%|0?[.,]\d{2,4}|1[.,]0+)", re.I)
RE_DETECCION = re.compile(r"anomal|incident|accident|abnormal|outlier|detect|crash|closure|disruption", re.I)
RE_OTRA_TAREA = re.compile(r"kinetics|ucf-?101|hmdb|action recognition|imputation|top-1|top-5|re-?identification|"
                           r"forecast(ing)? (accuracy|error)|prediction (accuracy|error)", re.I)
RE_SOLO_PRONOSTICO = re.compile(r"\b(mae|rmse|mape)\b", re.I)
RE_SURVEY = re.compile(r"\b(a survey|this survey|systematic (literature )?review|literature review|we review)\b", re.I)
RE_PREPRINT = re.compile(r"arxiv:\d{4}\.\d{4,5}|preprint submitted|under review", re.I)
RE_PROPUESTA = re.compile(r"(we propose|this (paper|article|study|work) (proposes|presents|introduces)|we present|we introduce|"
                          r"our (proposed )?(model|method|framework|approach))", re.I)


logging.getLogger("pypdf").setLevel(logging.ERROR)   # oculta avisos de fuentes que no afectan al texto
warnings.filterwarnings("ignore")


def texto_paginas(ruta):
    try:
        from pypdf import PdfReader
    except ImportError:
        sys.exit("Falta la librería pypdf. Instálala con:  pip install pypdf")
    try:
        lector = PdfReader(str(ruta))
        pags = []
        for p in lector.pages:
            t = p.extract_text() or ""
            t = unicodedata.normalize("NFKC", t)
            t = re.sub(r"-\s*\n\s*", "", t)
            pags.append(re.sub(r"\s+", " ", t))
        return pags
    except Exception as e:  # PDF dañado o protegido
        return [f"__ERROR__ {e}"]


def citas(pags, patron, maximo=3, contexto=110, filtro=None, desde=0):
    """Frases cortas (una por página) donde aparece el patrón, con su número de página."""
    out = []
    for n, t in enumerate(pags[desde:], desde + 1):
        for m in patron.finditer(t):
            ini, fin = max(0, m.start() - contexto), min(len(t), m.end() + contexto)
            frag = t[ini:fin].strip()
            if filtro and not filtro.search(frag):
                continue
            out.append(f'p. {n}: "...{frag}..."')
            break                                   # una cita por página
        if len(out) >= maximo:
            break
    return out


def metricas_de_deteccion(pags):
    """Citas donde una métrica lleva un valor numérico en un contexto de detección (no de otra tarea)."""
    out = []
    for n, t in enumerate(pags, 1):
        for m in RE_METRICA_NUM.finditer(t):
            ctx = t[max(0, m.start() - 220):m.end() + 220]
            if RE_DETECCION.search(ctx) and not RE_OTRA_TAREA.search(ctx):
                out.append(f'p. {n}: "...{t[max(0, m.start() - 110):m.end() + 60].strip()}..."')
                break
        if len(out) >= 3:
            break
    return out


def analizar(pags):
    """Devuelve (sugerencia, motivo, ev_metricas, ev_propuesta, ev_dl, ev_anomalia)."""
    if pags and pags[0].startswith("__ERROR__"):
        return "Revisar", "No se pudo leer el PDF (dañado o protegido): ábrelo manualmente", [], [], [], []
    total = " ".join(pags)
    if len(total) < 2000:
        return "Revisar", "Casi no hay texto extraíble (¿PDF escaneado?): léelo manualmente", [], [], [], []
    inicio = pags[0][:3000]
    ev_met = metricas_de_deteccion(pags)
    ev_prop = citas(pags[:3], RE_PROPUESTA, 2, contexto=160)
    ev_dl = citas(pags, RE_DL, 2)
    ev_an = citas(pags, RE_ANOMALIA, 2)
    if len(RE_SURVEY.findall(inicio)) >= 1 and not ev_met:
        return "Revisar (posible CE-5)", "El inicio sugiere una revisión o survey", ev_met, ev_prop, ev_dl, ev_an
    motivos = []
    if RE_PREPRINT.search(inicio):
        motivos.append("posible preprint (CE-5): verifica si está publicado")
    no_vial, vial = RE_NO_VIAL.findall(total), RE_VIAL.findall(total)
    if len(no_vial) > len(vial):
        motivos.append(f"predominan términos no viales ({len(no_vial)} vs {len(vial)}): posible CE-2")
    if len(RE_ANOMALIA.findall(total)) < 5:
        motivos.append("pocas menciones de anomalía/incidente: posible CE-3")
    if len(RE_DL.findall(total)) < 5:
        motivos.append("pocas menciones de Deep Learning: posible CE-4")
    if not ev_met:
        otra = RE_OTRA_TAREA.search(total)
        detalle = []
        if RE_SOLO_PRONOSTICO.search(total):
            detalle.append("solo aparecen errores de pronóstico (MAE/RMSE)")
        if otra:
            detalle.append(f"las métricas parecen de otra tarea («{otra.group(0)}»)")
        motivos.append("no se hallaron métricas de DETECCIÓN con valores" + (f" ({'; '.join(detalle)})" if detalle else "") +
                       ": posible CI-6 — revisa la sección de resultados")
    motivos.append("comprueba en «qué proponen» que el método propio sea Deep Learning (CE-4)")
    if len(motivos) > 1:
        return "Revisar", "; ".join(motivos), ev_met, ev_prop, ev_dl, ev_an
    return "Probable incluir", "Hay métricas de detección con valores, DL y anomalía de tráfico. " + motivos[0], ev_met, ev_prop, ev_dl, ev_an


def main():
    asegurar_utf8_consola()
    ruta = REVISION / "TC.xlsx"
    wb = load_workbook(ruta)
    ws = wb["Cribado"]
    cab = [c.value for c in ws[1]]
    col = lambda nombre: cab.index(nombre) + 1
    filas_csv, leidos, faltan = [], 0, []
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, 1).value
        if not rid:
            continue
        pdf = CRIBADO / "pdf" / f"{rid}.pdf"
        if not pdf.exists():
            faltan.append(rid)
            continue
        sug, motivo, met, prop, dl, an = analizar(texto_paginas(pdf))
        leidos += 1
        ws.cell(r, col("SUGERENCIA")).value = sug
        ws.cell(r, col("Motivo de la sugerencia")).value = motivo
        ws.cell(r, col("Evidencia sugerida: métricas de detección")).value = "\n".join(met) or "—"
        ws.cell(r, col("Evidencia sugerida: qué proponen")).value = "\n".join(prop) or "—"
        ws.cell(r, col("Evidencia sugerida: Deep Learning")).value = "\n".join(dl) or "—"
        ws.cell(r, col("Evidencia sugerida: anomalía / tráfico")).value = "\n".join(an) or "—"
        if not ws.cell(r, col(COL_TEXTO)).value:
            ws.cell(r, col(COL_TEXTO)).value = "Sí"
        ws.row_dimensions[r].height = 140
        filas_csv.append({"id": rid, "sugerencia": sug, "motivo": motivo, "metricas": " | ".join(met), "propuesta": " | ".join(prop),
                          "deep_learning": " | ".join(dl), "anomalia": " | ".join(an)})
    wb.save(ruta)
    escribir_csv(SALIDA / "05b_evidencias.csv", filas_csv, ["id", "sugerencia", "motivo", "metricas", "propuesta", "deep_learning", "anomalia"])
    print(f"PDF leídos: {leidos}. Sugerencias escritas en {ruta}.")
    if faltan:
        print(f"Sin PDF todavía ({len(faltan)}): {', '.join(faltan[:15])}{' ...' if len(faltan) > 15 else ''}")
        print("Descárgalos a cribado/pdf/<ID>.pdf y vuelve a ejecutar este script. Si no se consiguen, marca «No recuperado».")


if __name__ == "__main__":
    main()
