#!/usr/bin/env python3
"""PASO 10. Trae a TA.xlsx las decisiones que verificaste o cambiaste en Parsif.al.

Si revisas los estudios en Parsif.al (en vez de en Excel), al terminar:
  1. En Parsif.al: Conducting > Study Selection > All Sources > «Export Articles» (descarga articles.xls).
  2. Guarda el archivo en cribado/salida/parsifal/articles.xls.
  3. Ejecuta (desde la carpeta RSL):   python cribado/scripts/10_traer_de_parsifal.py
     (si falta la librería: pip install xlrd)

Correspondencia: Accepted → Incluir (o Duda si el comentario empieza por «[Duda»);
Rejected → Excluir con el criterio elegido en Parsif.al (o el que diga el comentario);
Duplicated → Excluir CE-1; Unclassified → sin decisión.
Actualiza DECISIÓN, CRITERIO y Comentario de TA.xlsx (no toca las columnas «… IA») y muestra
cuántas decisiones cambiaron respecto de lo que había en la hoja.
"""
import argparse
import re
import sys

from openpyxl import load_workbook

from comun import REVISION, SALIDA, asegurar_utf8_consola
from hojas import COL_CRITERIO, COL_DECISION

RE_CRIT = re.compile(r"\b(C[IE]-\d+)\b")
RE_ETIQUETA = re.compile(r"^\s*\[(Incluir|Duda|Excluir)[^\]]*\]\s*")


def leer_export(ruta):
    try:
        import xlrd
    except ImportError:
        sys.exit("Falta la librería xlrd. Instálala con:  pip install xlrd")
    ws = xlrd.open_workbook(str(ruta)).sheet_by_index(0)
    cab = [str(c.value) for c in ws.row(0)]
    filas = {}
    for i in range(1, ws.nrows):
        f = {cab[j]: str(ws.cell_value(i, j) or "").strip() for j in range(len(cab))}
        m = re.search(r"ID cribado (R\d+)", f.get("note", ""))  # Parsif.al puede dejar vacía la clave
        clave = m.group(1) if m else f.get("bibtex_key", "")
        if clave:
            filas[clave] = f
    return filas


def decision_de(f):
    estado, comentario = f.get("status", ""), f.get("comments", "")
    texto = RE_ETIQUETA.sub("", comentario)
    if estado == "Accepted":
        return ("Duda" if comentario.lstrip().startswith("[Duda") else "Incluir"), "", texto
    if estado == "Duplicated":
        return "Excluir", "CE-1", texto
    if estado == "Rejected":
        m = RE_CRIT.match(f.get("selection_criteria", "")) or RE_CRIT.search(comentario)
        return "Excluir", (m.group(1) if m else ""), texto
    return "", "", texto


def main():
    asegurar_utf8_consola()
    ap = argparse.ArgumentParser()
    ap.add_argument("archivo", nargs="?", default=str(SALIDA / "parsifal" / "articles.xls"))
    a = ap.parse_args()
    export = leer_export(a.archivo)
    ruta = REVISION / "TA.xlsx"
    wb = load_workbook(ruta)
    ws = wb["Cribado"]
    cab = [c.value for c in ws[1]]
    cd, cc, cm = cab.index(COL_DECISION) + 1, cab.index(COL_CRITERIO) + 1, cab.index("Comentario") + 1
    cambios, sin_dec, sin_crit, vistos = [], [], [], 0
    for fila in range(2, ws.max_row + 1):
        i = ws.cell(row=fila, column=1).value
        if not i or str(i) not in export:
            continue
        vistos += 1
        d, crit, com = decision_de(export[str(i)])
        if not d:
            sin_dec.append(i)
            continue
        if d == "Excluir" and not crit:
            sin_crit.append(i)
        antes = (ws.cell(row=fila, column=cd).value or "", ws.cell(row=fila, column=cc).value or "")
        if antes != (d, crit):
            cambios.append(f"{i}: {antes[0]} {antes[1]} → {d} {crit}".strip())
        ws.cell(row=fila, column=cd, value=d)
        ws.cell(row=fila, column=cc).value = crit or None
        ws.cell(row=fila, column=cm, value=com)
    wb.save(ruta)
    print(f"Estudios de Parsif.al encontrados en TA.xlsx: {vistos} de {ws.max_row - 1}.")
    print(f"Decisiones que cambiaron respecto de la hoja: {len(cambios)}")
    for c in cambios[:40]:
        print("   " + c)
    if sin_dec:
        print(f"Sin decidir en Parsif.al (Unclassified): {len(sin_dec)}. Primeros: {', '.join(map(str, sin_dec[:15]))}")
    if sin_crit:
        print(f"Rejected sin criterio: {len(sin_crit)}. Elige el criterio en Parsif.al. Primeros: {', '.join(map(str, sin_crit[:15]))}")
    print("Cuando todo esté decidido: python cribado/scripts/04_cerrar_etapa.py --etapa TA")


if __name__ == "__main__":
    main()
