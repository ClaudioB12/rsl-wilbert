#!/usr/bin/env python3
"""PASO 4. Une las hojas de los dos revisores, calcula el kappa y prepara la reunión de consenso.

Uso (desde la carpeta RSL), para --etapa piloto | TA | TC:

  A) Consolidar:   python cribado/scripts/04_consolidar.py --etapa TA
     - Lee cribado/revision/TA_R1.xlsx y TA_R2.xlsx.
     - Calcula el kappa de Cohen (Incluir/Duda = pasa; Excluir = no pasa).
     - Crea cribado/salida/consolidado_TA.xlsx: donde los revisores coinciden, la decisión
       final ya viene llenada; donde NO coinciden (en rojo) la deciden juntos o el tercer revisor.

  B) Cerrar:       python cribado/scripts/04_consolidar.py --etapa TA --cerrar
     - Verifica que todos los registros tengan decisión final (y criterio y evidencia si se excluyen).
     - Guarda el resultado en cribado/salida/consenso_TA.json y el kappa en salida/kappa.json.
"""
import argparse
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from comun import (CRITERIOS, REVISION, SALIDA, asegurar_utf8_consola, cargar_json, guardar_json,
                   interpretar_kappa, kappa_cohen, leer_csv)
from hojas import AJUSTE, AMARILLO, AZUL, BLANCO_NEGRITA, ROJO, _lista, leer_hoja_revisor


def pasa(decision):
    return {"Incluir": "Pasa", "Duda": "Pasa", "Excluir": "No pasa"}.get(decision, "")


def consolidar(etapa, forzar):
    r1 = leer_hoja_revisor(REVISION / f"{etapa}_R1.xlsx")
    r2 = leer_hoja_revisor(REVISION / f"{etapa}_R2.xlsx")
    destino = SALIDA / f"consolidado_{etapa}.xlsx"
    if destino.exists() and not forzar:
        sys.exit(f"Ya existe {destino}. Si ya empezaron la reunión de consenso, no lo regeneren "
                 f"(usa --forzar solo si quieres rehacerlo y perder lo escrito en él).")
    titulos = {r["id"]: r["titulo"] for r in leer_csv(SALIDA / "02_registros_unicos.csv")}

    vacios = [(i, rev) for rev, d in (("R1", r1), ("R2", r2)) for i, v in d.items() if not v["decision"]]
    vacios += [(i, rev) for rev, d in (("R1", r1), ("R2", r2)) for i, v in d.items()
               if v["decision"] == "Excluir" and not v["criterio"]]
    if vacios:
        print(f"ATENCIÓN: hay {len(vacios)} celdas incompletas (sin decisión, o «Excluir» sin criterio). Ejemplos:")
        for i, rev in vacios[:10]:
            print(f"   {rev} -> {i}")
        print("   Complétenlas en la hoja del revisor y vuelvan a ejecutar este paso.\n")

    piloto = set(cargar_json(SALIDA / "consenso_piloto.json", {}).keys()) if etapa == "TA" else set()
    es_tc = etapa == "TC"
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    cab = ["ID", "Título"]
    if es_tc:
        cab += ["R1 texto completo", "R2 texto completo"]
    cab += ["R1 decisión", "R1 criterio", "R1 evidencia / comentario", "R2 decisión", "R2 criterio",
            "R2 evidencia / comentario", "ACUERDO", "DECISIÓN FINAL", "CRITERIO FINAL"]
    if es_tc:
        cab += ["EVIDENCIA FINAL (obligatoria si se excluye)"]
    cab += ["Nota de la reunión de consenso"]
    ws.append(cab)
    anchos = {"ID": 8, "Título": 50, "ACUERDO": 16, "DECISIÓN FINAL": 14, "CRITERIO FINAL": 12}
    for i, c in enumerate(ws[1], 1):
        c.fill, c.font, c.alignment = AZUL, BLANCO_NEGRITA, Alignment(wrap_text=True)
        ws.column_dimensions[c.column_letter].width = anchos.get(c.value, 30 if "evidencia" in c.value.lower() else 12)
    pares, pares_tc, discrepancias = [], [], 0
    for i in sorted(set(r1) | set(r2)):
        a, b = r1.get(i, {}), r2.get(i, {})
        da, db = a.get("decision", ""), b.get("decision", "")
        final_d = final_c = final_e = ""
        if es_tc and (a.get("texto_completo") == "No recuperado" or b.get("texto_completo") == "No recuperado"):
            acuerdo = "SÍ" if a.get("texto_completo") == b.get("texto_completo") else "NO (recuperación)"
            if acuerdo == "SÍ":
                final_d = "No recuperado"
        elif not da or not db:
            acuerdo = "INCOMPLETO"
        elif pasa(da) == pasa(db):
            acuerdo = "SÍ"
            if pasa(da) == "Pasa":
                final_d = "Incluir"
            else:
                final_d = "Excluir"
                if a.get("criterio") == b.get("criterio"):
                    final_c = a.get("criterio", "")
                    if es_tc:
                        final_e = a.get("evidencia") or b.get("evidencia")
                else:
                    acuerdo = "SÍ (criterio distinto)"
        else:
            acuerdo = "NO"
        if not (es_tc and final_d == "No recuperado") and i not in piloto:
            pares.append((pasa(da), pasa(db)))
        if acuerdo.startswith("NO") or acuerdo in ("INCOMPLETO", "SÍ (criterio distinto)"):
            discrepancias += 1
        fila = [i, titulos.get(i, "")]
        if es_tc:
            fila += [a.get("texto_completo", ""), b.get("texto_completo", "")]
        fila += [da, a.get("criterio", ""), a.get("evidencia") or a.get("comentario", ""),
                 db, b.get("criterio", ""), b.get("evidencia") or b.get("comentario", ""),
                 acuerdo, final_d or None, final_c or None]
        if es_tc:
            fila += [final_e or None]
        fila += [None]
        ws.append(fila)
        for c in ws[ws.max_row]:
            c.alignment = AJUSTE
        if acuerdo != "SÍ":
            for c in ws[ws.max_row][cab.index("ACUERDO"):]:
                c.fill = ROJO if acuerdo.startswith("NO") or acuerdo == "INCOMPLETO" else AMARILLO
    n = ws.max_row
    L = lambda nombre: ws.cell(row=1, column=cab.index(nombre) + 1).column_letter
    _lista(ws, f"{L('DECISIÓN FINAL')}2:{L('DECISIÓN FINAL')}{n}",
           ["Incluir", "Excluir", "No recuperado"] if es_tc else ["Incluir", "Excluir"])
    _lista(ws, f"{L('CRITERIO FINAL')}2:{L('CRITERIO FINAL')}{n}", CRITERIOS)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(cab)).column_letter}{n}"

    k, nk, po = kappa_cohen(pares)
    ks = wb.create_sheet("Kappa", 0)
    ks.column_dimensions["A"].width = 46
    ks.column_dimensions["B"].width = 40
    filas = [("Etapa", {"piloto": "Piloto", "TA": "Título y resumen", "TC": "Texto completo"}[etapa]),
             ("Registros evaluados por ambos (n)", nk),
             ("Acuerdo observado", f"{po:.1%}" if po is not None else "-"),
             ("Kappa de Cohen", f"{k:.2f}" if k is not None else "-"),
             ("Interpretación (Landis y Koch)", interpretar_kappa(k)),
             ("Registros a discutir (en rojo o amarillo)", discrepancias)]
    if piloto:
        filas.append(("Registros del piloto excluidos de este kappa", len(piloto & (set(r1) | set(r2)))))
    for a_, b_ in filas:
        ks.append([a_, b_])
        ks.cell(row=ks.max_row, column=1).font = Font(bold=True)
    ks.append([])
    ks.append(["Siguiente paso", "Filtrar la columna ACUERDO por «NO» / «criterio distinto», discutir cada caso con el "
                                 "MANUAL_CRIBADO.md, escribir DECISIÓN FINAL, CRITERIO FINAL y la nota. Luego ejecutar "
                                 "04_consolidar.py --etapa " + etapa + " --cerrar"])
    ks.cell(row=ks.max_row, column=2).alignment = AJUSTE
    SALIDA.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    print(f"Kappa {etapa}: {k:.2f} ({interpretar_kappa(k)}), n = {nk}, acuerdo = {po:.1%}" if k is not None else "Kappa: sin datos")
    if etapa == "piloto" and k is not None and k < 0.6:
        print("El kappa del piloto es menor que 0,60: revisen juntos los desacuerdos, aclaren el MANUAL_CRIBADO.md "
              "(anoten el cambio) y repitan el piloto con otra muestra (03 ... --piloto 50 --semilla 2027 --forzar).")
    print(f"Registros a discutir: {discrepancias}.  Archivo: {destino}")


def cerrar(etapa):
    ruta = SALIDA / f"consolidado_{etapa}.xlsx"
    ws = load_workbook(ruta, data_only=True)["Consolidado"]
    cab = [c.value for c in ws[1]]
    ix = {c: i for i, c in enumerate(cab)}
    consenso, errores = {}, []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila[0]:
            continue
        i = str(fila[0])
        d = (fila[ix["DECISIÓN FINAL"]] or "").strip() if fila[ix["DECISIÓN FINAL"]] else ""
        c = (fila[ix["CRITERIO FINAL"]] or "").strip() if fila[ix["CRITERIO FINAL"]] else ""
        e = ""
        if "EVIDENCIA FINAL (obligatoria si se excluye)" in ix:
            e = fila[ix["EVIDENCIA FINAL (obligatoria si se excluye)"]] or ""
        nota = fila[ix["Nota de la reunión de consenso"]] or ""
        if not d:
            errores.append(f"{i}: falta DECISIÓN FINAL")
        elif d == "Excluir" and not c:
            errores.append(f"{i}: excluido sin CRITERIO FINAL")
        elif d == "Excluir" and etapa == "TC" and not str(e).strip():
            errores.append(f"{i}: excluido a texto completo sin EVIDENCIA FINAL")
        consenso[i] = {"decision": d, "criterio": c, "evidencia": str(e), "nota": str(nota),
                       "acuerdo": fila[ix["ACUERDO"]]}
    if errores:
        print(f"No se puede cerrar la etapa {etapa}: {len(errores)} problemas. Primeros:")
        for e in errores[:15]:
            print("   " + e)
        sys.exit(1)
    guardar_json(SALIDA / f"consenso_{etapa}.json", consenso)
    k_ws = load_workbook(ruta, data_only=True)["Kappa"]
    valores = {r[0]: r[1] for r in k_ws.iter_rows(values_only=True) if r and r[0]}
    kappas = cargar_json(SALIDA / "kappa.json", {})
    kappas[etapa] = {"kappa": valores.get("Kappa de Cohen"), "n": valores.get("Registros evaluados por ambos (n)"),
                     "acuerdo": valores.get("Acuerdo observado")}
    guardar_json(SALIDA / "kappa.json", kappas)
    from collections import Counter
    cuenta = Counter(v["decision"] for v in consenso.values())
    print(f"Etapa {etapa} cerrada: " + ", ".join(f"{k}: {v}" for k, v in cuenta.items()))
    if etapa == "TC":
        crit = Counter(v["criterio"] for v in consenso.values() if v["decision"] == "Excluir")
        print("Exclusiones por criterio: " + ", ".join(f"{k}: {v}" for k, v in sorted(crit.items())))


def main():
    asegurar_utf8_consola()
    ap = argparse.ArgumentParser()
    ap.add_argument("--etapa", required=True, choices=["piloto", "TA", "TC"])
    ap.add_argument("--cerrar", action="store_true")
    ap.add_argument("--forzar", action="store_true")
    args = ap.parse_args()
    cerrar(args.etapa) if args.cerrar else consolidar(args.etapa, args.forzar)


if __name__ == "__main__":
    main()
