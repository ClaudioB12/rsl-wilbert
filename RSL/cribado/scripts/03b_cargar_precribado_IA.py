#!/usr/bin/env python3
"""PASO 3b (opcional). Carga en TA.xlsx el pre-cribado propuesto por un asistente de IA.

El asistente leyó el título y el resumen de cada registro y aplicó el árbol de decisión del
MANUAL_CRIBADO.md. Su propuesta se guarda en dos sitios de la hoja:
  - columnas «Propuesta IA», «Criterio IA», «Confianza IA» y «Comentario IA» (no se tocan: son
    el registro de lo que propuso el asistente);
  - DECISIÓN, CRITERIO y Comentario, prellenadas con esa propuesta para que el autor la VERIFIQUE
    y la cambie donde no esté de acuerdo.
Al cerrar la etapa, 04_cerrar_etapa.py compara DECISIÓN con «Propuesta IA» y reporta cuántas
decisiones cambió el autor (acuerdo y kappa autor-IA), que se declaran en el artículo.

Uso (desde la carpeta RSL):
    python cribado/scripts/03b_cargar_precribado_IA.py cribado/salida/precribado_IA_TA.json
El JSON es una lista de {"id", "decision", "criterio", "comentario", "confianza"}.
No sobrescribe una DECISIÓN que ya tenga valor distinto de la propuesta por reglas.
"""
import argparse
import json
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from comun import REVISION, asegurar_utf8_consola
from hojas import COL_CRITERIO, COL_DECISION

COLS_IA = [("Propuesta IA", 11), ("Criterio IA", 9), ("Confianza IA", 10), ("Comentario IA", 50)]
ORDEN_CONF = {"baja": 0, "media": 1, "alta": 2}


def main():
    asegurar_utf8_consola()
    ap = argparse.ArgumentParser()
    ap.add_argument("json", help="archivo con las propuestas del asistente")
    a = ap.parse_args()
    propuestas = {p["id"]: p for p in json.load(open(a.json, encoding="utf-8"))}
    ruta = REVISION / "TA.xlsx"
    wb = load_workbook(ruta)
    ws = wb["Cribado"]
    cab = [c.value for c in ws[1]]
    for nombre, ancho in COLS_IA:
        if nombre not in cab:
            col = len(cab) + 1
            c = ws.cell(row=1, column=col, value=nombre)
            c.fill, c.font = PatternFill("solid", fgColor="7030A0"), Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[c.column_letter].width = ancho
            cab.append(nombre)
    idx = {n: cab.index(n) + 1 for n in cab if n}
    faltan, cargadas = [], 0
    for fila in range(2, ws.max_row + 1):
        i = ws.cell(row=fila, column=1).value
        if not i:
            continue
        p = propuestas.get(str(i))
        if not p:
            faltan.append(i)
            continue
        crit = p.get("criterio", "") if p["decision"] == "Excluir" else ""
        for nombre, valor in (("Propuesta IA", p["decision"]), ("Criterio IA", crit),
                              ("Confianza IA", p.get("confianza", "")), ("Comentario IA", p.get("comentario", ""))):
            ws.cell(row=fila, column=idx[nombre], value=valor).alignment = Alignment(wrap_text=True, vertical="top")
        actual = ws.cell(row=fila, column=idx[COL_DECISION]).value
        regla = (ws.cell(row=fila, column=idx["Comentario"]).value or "").startswith("Regla automática")
        if actual in (None, "") or regla:
            ws.cell(row=fila, column=idx[COL_DECISION], value=p["decision"])
            ws.cell(row=fila, column=idx[COL_CRITERIO], value=crit or None)
            ws.cell(row=fila, column=idx["Comentario"], value="IA (verificar): " + p.get("comentario", ""))
        cargadas += 1
    wb.save(ruta)
    print(f"Propuestas cargadas en {ruta}: {cargadas}")
    if faltan:
        print(f"Sin propuesta ({len(faltan)}): {', '.join(map(str, faltan[:20]))}")
        sys.exit(1)
    print("Ahora verifica cada fila (empieza por «Confianza IA» = baja y por las exclusiones) y corrige "
          "DECISIÓN/CRITERIO donde no estés de acuerdo. No edites las columnas «… IA».")


if __name__ == "__main__":
    main()
