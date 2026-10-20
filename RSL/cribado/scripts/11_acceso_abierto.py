#!/usr/bin/env python3
"""PASO 11. Clasifica cada registro único según tenga o no el texto completo en ACCESO ABIERTO,
y aplica esa condición (parte del criterio CI-2) en la hoja de título y resumen.

Por qué: el protocolo exige que el texto completo esté en acceso abierto (CI-2). En la búsqueda,
el filtro «Open Access» se usó solo en algunas bases (Scopus, IEEE) y no en otras (WoS,
ScienceDirect, ACM). Para que la regla sea IGUAL en todas, se verifica registro por registro aquí.

Definición: acceso abierto = el texto completo se puede leer legalmente sin suscripción
(gold, diamond, hybrid, bronze o green), según OpenAlex (índice abierto que integra Unpaywall)
o según el indicador de acceso abierto de la propia base. La fecha de verificación se guarda.

Fuentes del dato, en este orden:
  1. OpenAlex (https://api.openalex.org), por DOI; si no hay DOI, por título exacto.
  2. Indicador de la base exportada: Scopus («All Open Access» en note), Web of Science (campo OA)
     e IEEE Xplore (exportación con el filtro «Open Access Only», si se guardó en
     exportaciones/IEEE/acceso_abierto/).
  Es de acceso abierto si CUALQUIERA de las fuentes lo indica. Si ninguna fuente tiene datos del
  registro (sin DOI y sin coincidencia de título), queda «Por verificar» para revisarlo a mano.

Uso (desde la carpeta RSL; necesita internet, tarda ~5 minutos; si se corta, vuelve a ejecutarlo
y continúa donde quedó):
    python cribado/scripts/11_acceso_abierto.py                 # clasifica y muestra el resumen
    python cribado/scripts/11_acceso_abierto.py --aplicar       # además lo aplica en TA.xlsx
    python cribado/scripts/11_acceso_abierto.py --sin-internet  # solo con los datos de las bases

Salidas:
  cribado/salida/11_acceso_abierto.csv   ID, estado (Sí / No / Por verificar), tipo, enlace al
                                         texto completo gratuito (sirve para descargar los PDF).
  Con --aplicar, en TA.xlsx: los registros que NO son de acceso abierto pasan a
  «Excluir — CI-2 (no es de acceso abierto)», salvo los ya excluidos por un criterio anterior del
  árbol (CE-1, CE-5, CI-3), que conservan el suyo. Las columnas «… IA» no se tocan.
"""
import argparse
import csv
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date

from openpyxl import load_workbook

from comun import (EXPORTACIONES, REVISION, SALIDA, _campos_bibtex, asegurar_utf8_consola, escribir_csv, leer_csv,
                   leer_texto, normalizar_doi, normalizar_titulo)
from hojas import COL_CRITERIO, COL_DECISION

API = "https://api.openalex.org/works"
CACHE = SALIDA / "11_openalex_cache.json"
CRITERIOS_ANTERIORES = {"CE-1", "CE-5", "CI-3"}  # en el árbol van antes que CI-2


def entradas_bib(ruta):
    """(campos) de cada entrada de un .bib, con los nombres de campo en minúsculas."""
    texto = leer_texto(ruta)
    for m in re.finditer(r"@(\w+)\s*\{", texto):
        if m.group(1).lower() in ("comment", "preamble", "string"):
            continue
        prof, j = 1, m.end()
        while j < len(texto) and prof:
            prof += {"{": 1, "}": -1}.get(texto[j], 0)
            j += 1
        cuerpo = texto[m.end():j - 1]
        yield _campos_bibtex(cuerpo.split(",", 1)[1] if "," in cuerpo else "")


def oa_de_las_bases():
    """Devuelve (marcas, presentes).

    marcas: {clave: {"Scopus: Gold Open Access", ...}} de los registros que la base marca como abiertos.
    presentes: {clave: {"Scopus", ...}} de las bases que tienen ese registro y SÍ informan el acceso
    abierto (si una de ellas lo tiene y no lo marca, esa base dice que no es abierto).
    clave = DOI normalizado o título normalizado.
    """
    marcas, presentes = {}, {}

    def claves(c):
        return [k for k in (normalizar_doi(c.get("doi", "")), normalizar_titulo(c.get("title", ""))) if k]

    def marcar(c, texto):
        for k in claves(c):
            marcas.setdefault(k, set()).add(texto)

    def presente(c, base):
        for k in claves(c):
            presentes.setdefault(k, set()).add(base)

    for ruta in (EXPORTACIONES / "Scopus").glob("*.bib"):
        for c in entradas_bib(ruta):
            presente(c, "Scopus")
            tipos = re.findall(r"(?:All|Gold|Green|Hybrid Gold|Bronze)\s+Open Access", c.get("note", ""), re.I)
            if tipos:
                marcar(c, "Scopus: " + ", ".join(dict.fromkeys(tipos)))
    for ruta in (EXPORTACIONES / "WoS").glob("*.bib"):
        for c in entradas_bib(ruta):
            presente(c, "WoS")
            if c.get("oa", "").strip():
                marcar(c, "WoS: " + c["oa"].strip())
    ieee_oa = list((EXPORTACIONES / "IEEE" / "acceso_abierto").glob("*.bib"))
    for ruta in ieee_oa:
        for c in entradas_bib(ruta):
            marcar(c, "IEEE: Open Access Only")
    if ieee_oa:  # solo si se guardó la exportación filtrada, la ausencia en ella significa «no abierto»
        for ruta in (EXPORTACIONES / "IEEE").glob("*.bib"):
            for c in entradas_bib(ruta):
                presente(c, "IEEE")
    return marcas, presentes


class LimiteDiario(ConnectionError):
    pass


def pedir(url):
    """JSON de la URL; None si OpenAlex responde que la consulta no tiene respuesta (404, 400…).

    Si OpenAlex pide esperar (429) o falla de forma temporal (5xx, red), reintenta con pausas
    crecientes (hasta ~2 minutos) antes de rendirse.
    """
    req = urllib.request.Request(url, headers={"User-Agent": "RSL-UPeU/1.0 (cribado de acceso abierto)"})
    ultimo = ""
    for intento in range(7):
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                return json.load(r)
        except urllib.error.HTTPError as e:
            if 400 <= e.code < 500 and e.code != 429:
                return None
            ultimo = f"HTTP {e.code}"
            espera = max(int(e.headers.get("Retry-After") or 0) if e.headers else 0, 2 ** intento)
            if espera > 300:  # límite diario de OpenAlex agotado: no tiene sentido esperar horas
                hora = time.strftime("%H:%M", time.localtime(time.time() + espera))
                raise LimiteDiario(f"OpenAlex alcanzó su límite diario de consultas; se puede continuar después de las {hora}")
            print(f"    (OpenAlex respondió {ultimo}; reintento {intento + 1} de 7 en {espera} s)", flush=True)
            time.sleep(espera)
        except Exception as e:
            ultimo = type(e).__name__
            print(f"    (sin respuesta: {ultimo}; reintento {intento + 1} de 7 en {2 ** intento} s)", flush=True)
            time.sleep(2 ** intento)
    raise ConnectionError(f"{ultimo} al consultar {url}")


def consultar_openalex(r, correo):
    """Devuelve {"oa": bool, "tipo": str, "url": str, "por": "doi"|"titulo"} o None si no se encontró."""
    extra = f"&mailto={urllib.parse.quote(correo)}" if correo else ""
    doi = normalizar_doi(r["doi"])
    if doi:
        w = pedir(f"{API}/doi:{urllib.parse.quote(doi)}?select=id,title,open_access{extra}")
        if w:
            oa = w.get("open_access") or {}
            return {"oa": bool(oa.get("is_oa")), "tipo": oa.get("oa_status") or "", "url": oa.get("oa_url") or "", "por": "doi"}
    titulo = re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", r["titulo"])).strip()
    if titulo and not re.match(r"proceedings of\b", titulo, re.I):  # volúmenes de actas: no son artículos (CE-5)
        q = urllib.parse.quote(f"title.search:{titulo}")
        res = pedir(f"{API}?filter={q}&per-page=5&select=id,title,open_access{extra}") or {}
        for w in res.get("results", []):
            if normalizar_titulo(w.get("title") or "") == normalizar_titulo(r["titulo"]):
                oa = w.get("open_access") or {}
                return {"oa": bool(oa.get("is_oa")), "tipo": oa.get("oa_status") or "", "url": oa.get("oa_url") or "", "por": "titulo"}
    return None


def aplicar_en_TA(estado):
    ruta = REVISION / "TA.xlsx"
    wb = load_workbook(ruta)
    ws = wb["Cribado"]
    cab = [c.value for c in ws[1]]
    cd, cc, cm = cab.index(COL_DECISION) + 1, cab.index(COL_CRITERIO) + 1, cab.index("Comentario") + 1
    if "Acceso abierto" not in cab:
        ws.cell(row=1, column=len(cab) + 1, value="Acceso abierto")
        cab.append("Acceso abierto")
    ca = cab.index("Acceso abierto") + 1
    cambiados, revertidos = 0, []
    for fila in range(2, ws.max_row + 1):
        i = ws.cell(row=fila, column=1).value
        e = estado.get(str(i))
        if not e:
            continue
        ws.cell(row=fila, column=ca, value=f"{e['acceso_abierto']} ({e['detalle']})")
        crit_actual = ws.cell(row=fila, column=cc).value or ""
        dec_actual = ws.cell(row=fila, column=cd).value or ""
        com_actual = ws.cell(row=fila, column=cm).value or ""
        # Si una corrida anterior lo excluyó por CI-2 con datos parciales y ahora consta como abierto,
        # se devuelve la propuesta original (columnas «… IA»).
        if e["acceso_abierto"] != "No" and crit_actual == "CI-2" and com_actual.startswith("No es de acceso abierto") \
                and "Propuesta IA" in cab:
            prop = ws.cell(row=fila, column=cab.index("Propuesta IA") + 1).value or ""
            crit_ia = ws.cell(row=fila, column=cab.index("Criterio IA") + 1).value or ""
            com_ia = ws.cell(row=fila, column=cab.index("Comentario IA") + 1).value or ""
            ws.cell(row=fila, column=cd).value = prop or None
            ws.cell(row=fila, column=cc).value = (crit_ia if prop == "Excluir" else None) or None
            ws.cell(row=fila, column=cm).value = f"IA (verificar): {com_ia}" if com_ia else None
            revertidos.append(str(i))
            continue
        if e["acceso_abierto"] == "No" and not (dec_actual == "Excluir" and crit_actual in CRITERIOS_ANTERIORES):
            if (dec_actual, crit_actual) != ("Excluir", "CI-2"):
                cambiados += 1
            antes = f" Antes: {dec_actual} {crit_actual}".rstrip() + "." if dec_actual else ""
            ws.cell(row=fila, column=cd, value="Excluir")
            ws.cell(row=fila, column=cc, value="CI-2")
            ws.cell(row=fila, column=cm, value=f"No es de acceso abierto ({e['detalle']}) → CI-2.{antes}")
    wb.save(ruta)
    if revertidos:
        print(f"Vueltos a la propuesta original (ahora constan como de acceso abierto): {', '.join(revertidos)}")
    return cambiados


def main():
    asegurar_utf8_consola()
    ap = argparse.ArgumentParser()
    ap.add_argument("--aplicar", action="store_true", help="aplicar el resultado en revision/TA.xlsx")
    ap.add_argument("--sin-internet", action="store_true", help="usar solo los indicadores de las bases")
    ap.add_argument("--correo", default="", help="tu correo (OpenAlex responde más rápido si lo das)")
    a = ap.parse_args()

    registros = leer_csv(SALIDA / "02_registros_unicos.csv")
    marcas, presentes = oa_de_las_bases()
    cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}
    hoy = date.today().isoformat()

    if not a.sin_internet:
        # Primero los que siguen en juego (Incluir/Duda): para ellos el acceso abierto decide; los ya
        # excluidos por otro criterio van al final.
        try:
            from hojas import leer_hoja
            dec = {i: v["decision"] for i, v in leer_hoja(REVISION / "TA.xlsx").items()}
        except Exception:
            dec = {}
        pendientes = sorted((r for r in registros if r["id"] not in cache), key=lambda r: dec.get(r["id"]) == "Excluir")
        en_juego = sum(1 for r in pendientes if dec.get(r["id"]) != "Excluir")
        print(f"  De ellos, {en_juego} están en Incluir/Duda (se consultan primero).")
        print(f"Consultando OpenAlex: {len(pendientes)} registros (ya consultados: {len(cache)}).")
        try:
            for n, r in enumerate(pendientes, 1):
                try:
                    cache[r["id"]] = consultar_openalex(r, a.correo) or {"no_encontrado": True}
                    cache[r["id"]]["fecha"] = hoy
                except LimiteDiario as e:
                    print(f"\n{e}. Lo consultado se guardó.")
                    break
                except ConnectionError as e:
                    print(f"\nSin conexión ({e}). Lo consultado se guardó; vuelve a ejecutar el script para continuar.")
                    break
                if n % 10 == 0:
                    CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
                if n % 50 == 0:
                    print(f"  … {n} de {len(pendientes)}", flush=True)
                time.sleep(0.12)
        except KeyboardInterrupt:
            CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
            sys.exit("\nDetenido con Ctrl+C. Lo consultado se guardó; vuelve a ejecutar el script para continuar.")
        CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")

    filas, estado = [], {}
    for r in registros:
        claves = [k for k in (normalizar_doi(r["doi"]), normalizar_titulo(r["titulo"])) if k]
        de_base = sorted({m for k in claves for m in marcas.get(k, ())})
        bases_con_dato = sorted({b for k in claves for b in presentes.get(k, ())})
        oa = cache.get(r["id"], {}) if not a.sin_internet else {}
        fuentes = []
        if oa.get("oa"):
            fuentes.append(f"OpenAlex: {oa.get('tipo') or 'oa'}")
        fuentes += de_base
        if fuentes:
            valor = "Sí"
        elif "oa" in oa or bases_con_dato:
            valor = "No"
            if "oa" in oa:
                fuentes.append(f"OpenAlex: {oa.get('tipo') or 'closed'}")
            fuentes += [f"{b}: sin acceso abierto" for b in bases_con_dato]
        else:
            valor = "Por verificar"
            fuentes.append("sin datos: verificar a mano (sitio de la editorial o Google Scholar)")
        detalle = "; ".join(fuentes) + (f"; verificado {oa['fecha']}" if oa.get("fecha") else "")
        estado[r["id"]] = {"acceso_abierto": valor, "detalle": detalle}
        filas.append({"id": r["id"], "titulo": r["titulo"], "anio": r["anio"], "doi": r["doi"],
                      "base_atribuida": r["base_atribuida"], "acceso_abierto": valor, "detalle": detalle,
                      "enlace_texto_completo": oa.get("url", "")})
    escribir_csv(SALIDA / "11_acceso_abierto.csv", filas,
                 ["id", "titulo", "anio", "doi", "base_atribuida", "acceso_abierto", "detalle", "enlace_texto_completo"])

    from collections import Counter
    print("\nAcceso abierto por base (registros únicos):")
    for b in ["Scopus", "WoS", "IEEE", "ScienceDirect", "ACM"]:
        c = Counter(f["acceso_abierto"] for f in filas if f["base_atribuida"] == b)
        print(f"  {b:<14} Sí={c['Sí']:>4}  No={c['No']:>4}  Por verificar={c['Por verificar']:>4}")
    c = Counter(f["acceso_abierto"] for f in filas)
    print(f"  {'TOTAL':<14} Sí={c['Sí']:>4}  No={c['No']:>4}  Por verificar={c['Por verificar']:>4}")
    print(f"Detalle: {SALIDA / '11_acceso_abierto.csv'}")
    if a.aplicar:
        n = aplicar_en_TA(estado)
        print(f"TA.xlsx actualizado: {n} registros pasaron a «Excluir — CI-2 (no es de acceso abierto)».")
        print("Los «Por verificar» no se tocaron: revísalos a mano (columna «Acceso abierto» de TA.xlsx).")
    elif c["No"]:
        print("Para aplicarlo en TA.xlsx, vuelve a ejecutar con --aplicar.")


if __name__ == "__main__":
    main()
